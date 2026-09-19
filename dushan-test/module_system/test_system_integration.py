from uuid import uuid4

import pytest
from httpx import ASGITransport, AsyncClient

from framework.starter_security.core.security_service import SecurityService
from framework.starter_security.definitions.constants.security_error_codes import (
    SecurityErrorCodes,
)
from framework.starter_security.definitions.enums.security_realm import SecurityRealm
from framework.starter_web.routing.route_policy import RoutePolicy
from module_system.service.notification.bo.notice_message_create_bo import NoticeMessageCreateBO

pytestmark = pytest.mark.asyncio(loop_scope="module")


async def test_job_workload_and_persistent_delivery_claim(admin_client, system_app):
    from datetime import datetime, timedelta, timezone

    from framework.starter_job.definitions.enums.job_trigger_kind import JobTriggerKind
    from framework.starter_job.model.job_context import JobContext
    from framework.starter_mq.exception.message_result_unknown import MessageResultUnknown
    from module_system.dal.dataobject.announcement.announcement_do import (
        AnnouncementDO,
    )
    from module_system.dal.dataobject.mail.mail_log_do import MailLogDO
    from module_system.dal.dataobject.notification.notice_message_do import NoticeMessageDO
    from module_system.dal.mapper.announcement.announcement_mapper import (
        AnnouncementMapper,
    )
    from module_system.dal.mapper.mail.mail_log_mapper import MailLogMapper
    from module_system.dal.mapper.notification.notice_message_mapper import NoticeMessageMapper
    from module_system.framework.notification.delivery.delivery_definite_failure import (
        DeliveryDefiniteFailure,
    )
    from module_system.job.announcement.announcement_publish_job import AnnouncementPublishJob
    from module_system.job.system_job_parameters import SystemJobParameters
    from module_system.service.auth.system_workload_service import SystemWorkloadService
    from module_system.service.notification.notification_delivery_service import (
        NotificationDeliveryService,
    )

    application = system_app.state.application_context
    with application.execution(), system_app.state.database.scope():
        security = application.container.get(SecurityService)
        mapper = application.container.get(AnnouncementMapper)
        async with security.authorized(
            admin_client.headers["Authorization"].removeprefix("Bearer "),
            RoutePolicy(tenant_required=True, realm=SecurityRealm.TENANT),
        ):
            now = datetime.now(timezone.utc).replace(tzinfo=None)
            announcement = await mapper.insert(
                AnnouncementDO(
                    title="Scheduled" + uuid4().hex,
                    content="Test",
                    status=1,
                    category=1,
                    publisher="test",
                    publish_time=now - timedelta(minutes=1),
                )
            )
            log = await application.container.get(MailLogMapper).insert(
                MailLogDO(
                    user_type=2,
                    to_mail="receiver@example.com",
                    account_id=1,
                    from_mail="sender@example.com",
                    template_id=1,
                    template_code="test",
                    template_title="Test",
                    template_content="Test",
                    template_params={},
                )
            )
        context = JobContext(
            job_id="1",
            handler_key="system.announcement.publish",
            request_id=uuid4().hex,
            attempt=1,
            trigger=JobTriggerKind.MANUAL,
            scheduled_at=datetime.now(timezone.utc),
            tenant_id="1",
        )
        async with application.container.get(SystemWorkloadService).scope(
            "system.announcement.publish", "1"
        ):
            await application.container.get(AnnouncementPublishJob).execute(
                SystemJobParameters(), context
            )
            assert (await mapper.select_by_id(announcement.id)).status == 2
            messages = application.container.get(NoticeMessageMapper)
            count = await messages.count(NoticeMessageDO.notice_title == announcement.title)
            assert count > 0
            await application.container.get(AnnouncementPublishJob).execute(
                SystemJobParameters(), context
            )
            assert await messages.count(NoticeMessageDO.notice_title == announcement.title) == count
        async with application.container.get(SystemWorkloadService).scope("system.mail.send", "1"):
            delivery = application.container.get(NotificationDeliveryService)
            attempt = await delivery.claim("mail", log.id)
            assert attempt is not None
            with pytest.raises(DeliveryDefiniteFailure):
                await delivery.claim("mail", log.id)
            await delivery.started("mail", log.id, attempt)
            with pytest.raises(MessageResultUnknown):
                await delivery.claim("mail", log.id)
            await delivery.finish(
                "mail", log.id, attempt, send_status=10, send_message_id="local-result"
            )
            assert await delivery.claim("mail", log.id) is None


async def test_registration_and_logout_revoke_cookie_and_session(system_app):
    async with AsyncClient(
        transport=ASGITransport(app=system_app), base_url="http://testserver"
    ) as client:
        payload = {
            "username": "r" + uuid4().hex[:10],
            "nickname": "Registered",
            "password": "Password123",
        }
        response = (await client.post("/admin-api/system/auth/register", json=payload)).json()
        assert response["code"] == 0, response
        assert "refreshToken" not in response["data"]
        client.headers["Authorization"] = "Bearer " + response["data"]["accessToken"]
        assert (await client.get("/admin-api/system/auth/codes")).json()["code"] == 0
        result = (
            await client.post(
                "/admin-api/system/auth/logout", headers={"Origin": "http://testserver"}
            )
        ).json()
        assert result["code"] == 0, result
        assert client.cookies.get("system_refresh") is None
        assert (await client.get("/admin-api/system/auth/codes")).json()[
            "code"
        ] == SecurityErrorCodes.REVOKED.code
        assert (await client.post("/admin-api/system/auth/register", json=payload)).json()[
            "code"
        ] != 0


async def test_login_and_current_user(system_app):
    async with AsyncClient(
        transport=ASGITransport(app=system_app), base_url="http://testserver"
    ) as client:
        response = await client.post(
            "/admin-api/system/auth/login", json={"username": "admin", "password": "admin123"}
        )
        body = response.json()
        assert body["code"] == 0, body
        assert "refreshToken" not in body["data"]
        assert client.cookies.get("system_refresh")
        client.headers["Authorization"] = "Bearer " + body["data"]["accessToken"]
        application = system_app.state.application_context
        with application.execution(), system_app.state.database.scope():
            async with application.container.get(SecurityService).authorized(
                body["data"]["accessToken"],
                RoutePolicy(tenant_required=True, realm=SecurityRealm.TENANT),
            ):
                assert application.container.get(SecurityService).context.require().account_id
        response = await client.get("/admin-api/system/auth/get-permission-info")
        assert response.json()["code"] == 0, response.json()
        assert response.json()["data"]["user"]["username"] == "admin"


async def test_batch_announcement_delete_uses_repeated_ids(admin_client, system_database):
    ids = []
    for _ in range(3):
        created = (
            await admin_client.post(
                "/admin-api/system/announcement/create",
                json={
                    "title": "Batch" + uuid4().hex,
                    "content": "Test",
                    "publisher": "test",
                    "status": 0,
                    "category": 1,
                },
            )
        ).json()
        assert created["code"] == 0, created
        ids.append(created["data"])
    for params in ({}, {"ids": ""}, {"ids": ",".join(ids)}, [("ids", ids[0]), ("ids", "invalid")]):
        rejected = (
            await admin_client.delete("/admin-api/system/announcement/delete-list", params=params)
        ).json()
        assert rejected["code"] == 422, rejected
    with system_database[2].cursor() as cursor:
        cursor.execute(
            "SELECT id FROM system_announcement WHERE id IN (%s,%s,%s) AND deleted=0", ids
        )
        assert {str(row[0]) for row in cursor.fetchall()} == set(ids)
    deleted = (
        await admin_client.delete(
            "/admin-api/system/announcement/delete-list",
            params=[("ids", value) for value in ids[:2]],
        )
    ).json()
    assert deleted["code"] == 0 and deleted["data"] == 2, deleted
    with system_database[2].cursor() as cursor:
        cursor.execute(
            "SELECT id FROM system_announcement WHERE id IN (%s,%s,%s) AND deleted=0", ids
        )
        assert [str(row[0]) for row in cursor.fetchall()] == ids[2:]
    single = (
        await admin_client.delete(
            "/admin-api/system/announcement/delete-list", params={"ids": ids[2]}
        )
    ).json()
    assert single["code"] == 0 and single["data"] == 1, single


async def test_notice_read_requires_json_string_array(admin_client):
    path = "/admin-api/system/notification/message/update-read"
    for ids in ([], "1,2", [1], ["1,2"]):
        rejected = (await admin_client.put(path, json={"ids": ids})).json()
        assert rejected["code"] == 422, rejected
    accepted = (await admin_client.put(path, json={"ids": [str((1 << 63) - 1)]})).json()
    assert accepted["code"] == 0 and accepted["data"] is True, accepted


async def test_announcement_publisher_contract_rejects_before_database_write(
    admin_client, system_database
):
    payload = {"title": "Publisher" + uuid4().hex, "content": "Test", "status": 0, "category": 1}
    for fields in ({}, {"publisher": None}, {"publisher": ""}, {"publisher": "x" * 65}):
        response = (
            await admin_client.post(
                "/admin-api/system/announcement/create", json={**payload, **fields}
            )
        ).json()
        assert response["code"] == 422, response
    with system_database[2].cursor() as cursor:
        cursor.execute(
            "SELECT COUNT(*) FROM system_announcement WHERE title=%s", (payload["title"],)
        )
        assert cursor.fetchone()[0] == 0
    created = (
        await admin_client.post(
            "/admin-api/system/announcement/create", json={**payload, "publisher": "Editor"}
        )
    ).json()
    assert created["code"] == 0, created
    changed = (
        await admin_client.put(
            "/admin-api/system/announcement/update",
            json={**payload, "id": created["data"], "publisher": "Updated"},
        )
    ).json()
    assert changed["code"] == 0, changed
    rejected = (
        await admin_client.put(
            "/admin-api/system/announcement/update",
            json={**payload, "id": created["data"], "publisher": None},
        )
    ).json()
    assert rejected["code"] == 422, rejected
    found = (
        await admin_client.get("/admin-api/system/announcement/get", params={"id": created["data"]})
    ).json()
    assert found["code"] == 0 and found["data"]["publisher"] == "Updated", found


async def test_sms_channel_cannot_change_provider_but_can_rotate_key(admin_client):
    from module_system.definitions.constants.error_code_constants import ErrorCodeConstants

    payload = {
        "signature": "Provider",
        "code": "ALIYUN",
        "status": 0,
        "apiKey": "test-key",
        "apiSecret": "first",
    }
    invalid = (
        await admin_client.post(
            "/admin-api/system/sms/channel/create", json={**payload, "signature": "x" * 13}
        )
    ).json()
    assert invalid["code"] == 422, invalid
    created = (await admin_client.post("/admin-api/system/sms/channel/create", json=payload)).json()
    assert created["code"] == 0, created
    rejected = (
        await admin_client.put(
            "/admin-api/system/sms/channel/update",
            json={**payload, "id": created["data"], "code": "TENCENT", "apiKey": "key sdk-app"},
        )
    ).json()
    assert rejected["code"] == ErrorCodeConstants.SMS_CHANNEL_CODE_IMMUTABLE.code, rejected
    updated = (
        await admin_client.put(
            "/admin-api/system/sms/channel/update",
            json={**payload, "id": created["data"], "apiSecret": "rotated"},
        )
    ).json()
    assert updated["code"] == 0, updated
    found = (
        await admin_client.get("/admin-api/system/sms/channel/get", params={"id": created["data"]})
    ).json()
    assert found["code"] == 0 and found["data"]["code"] == "ALIYUN", found


async def test_mail_controller_camel_input_reaches_structured_log_command(
    admin_client, system_database
):
    from module_system.definitions.enums.mail.mail_send_status_enum import MailSendStatusEnum

    suffix = uuid4().hex[:10]
    account = (
        await admin_client.post(
            "/admin-api/system/mail/account/create",
            json={
                "mail": "sender@example.com",
                "username": "sender",
                "password": "test-only",
                "host": "smtp.example.com",
                "port": 465,
                "sslEnable": True,
                "starttlsEnable": False,
            },
        )
    ).json()
    assert account["code"] == 0, account
    template = (
        await admin_client.post(
            "/admin-api/system/mail/template/create",
            json={
                "name": "Contract",
                "code": "contract_" + suffix,
                "accountId": account["data"],
                "nickname": "Contract",
                "title": "Hello",
                "content": "Hello {name}",
                "status": 0,
            },
        )
    ).json()
    assert template["code"] == 0, template
    sent = (
        await admin_client.post(
            "/admin-api/system/mail/template/send-mail",
            json={
                "toMails": ["recipient@example.com"],
                "ccMails": ["cc@example.com"],
                "bccMails": ["bcc@example.com"],
                "templateCode": "contract_" + suffix,
                "templateParams": {"name": "Contract"},
            },
        )
    ).json()
    assert sent["code"] == 0, sent
    assert isinstance(sent["data"], str)
    with system_database[2].cursor() as cursor:
        cursor.execute(
            "SELECT to_mail,cc_mail,bcc_mail,template_content,send_status FROM system_mail_log WHERE id=%s",
            (int(sent["data"]),),
        )
        assert cursor.fetchone() == (
            "recipient@example.com",
            "cc@example.com",
            "bcc@example.com",
            "Hello Contract",
            MailSendStatusEnum.IGNORE.code,
        )


async def test_export_query_keeps_selected_fields_with_filters(admin_client):
    from io import BytesIO

    from openpyxl import load_workbook

    exported = await admin_client.get(
        "/admin-api/system/mail/template/export-excel",
        params={"fields": "name", "pageSize": 5, "status": 0},
    )
    assert exported.content.startswith(b"PK"), exported.text[:200]
    workbook = load_workbook(BytesIO(exported.content), read_only=True)
    try:
        assert len(next(workbook.active.iter_rows(values_only=True))) == 1
    finally:
        workbook.close()


async def test_social_binding_uses_server_identity(admin_client, monkeypatch):
    from module_system.definitions.enums.social.social_type_enum import SocialTypeEnum
    from module_system.service.social.social_user_service_impl import SocialUserServiceImpl

    captured = []

    async def bind(self, req):
        captured.append(req)
        return "contract-openid"

    monkeypatch.setattr(SocialUserServiceImpl, "bind_social_user", bind)
    payload = {"type": next(iter(SocialTypeEnum)).code, "code": "test-code", "state": "state"}
    result = (await admin_client.post("/admin-api/system/social/user/bind", json=payload)).json()
    assert result["code"] == 0, result
    assert captured[0].user_id > 0 and captured[0].user_type == 2
    invalid = (
        await admin_client.post(
            "/admin-api/system/social/user/bind", json={**payload, "userId": "999"}
        )
    ).json()
    assert invalid["code"] == 422, invalid
    assert len(captured) == 1


async def test_user_import_form_vo_preserves_camel_case_flag(admin_client, monkeypatch):
    from io import BytesIO

    from openpyxl import Workbook

    from module_system.controller.admin.user.vo.user.user_import_resp_vo import UserImportRespVO
    from module_system.service.user.admin_user_service_impl import AdminUserServiceImpl

    captured = []

    async def receive_import(self, rows, is_update_support):
        captured.append((rows, is_update_support))
        return UserImportRespVO(create_usernames=[], update_usernames=[], failure_usernames={})

    monkeypatch.setattr(AdminUserServiceImpl, "import_user_list", receive_import)
    workbook = Workbook()
    workbook.active.append(["用户账号", "用户昵称"])
    workbook.active.append(["contract-user", "Contract User"])
    content = BytesIO()
    workbook.save(content)
    workbook.close()
    result = (
        await admin_client.post(
            "/admin-api/system/user/import",
            data={"updateSupport": "true"},
            files={
                "file": (
                    "users.xlsx",
                    content.getvalue(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
    ).json()
    assert result["code"] == 0, result
    assert captured[0][1] is True
    assert captured[0][0][0].username == "contract-user"


async def test_online_device_path_vo_and_verified_session(admin_client):
    from module_system.definitions.constants.error_code_constants import ErrorCodeConstants

    result = (await admin_client.get("/admin-api/system/user/profile/online-devices")).json()
    assert result["code"] == 0, result
    assert any(device["isCurrent"] for device in result["data"])
    denied = (await admin_client.delete("/admin-api/system/user/profile/online-devices/1")).json()
    assert denied["code"] == ErrorCodeConstants.USER_KICKOUT_DEVICE_NOT_OWNER.code, denied


async def test_sms_controller_structured_command_persists_disabled_template_log(
    admin_client, system_app, system_database
):
    from module_system.dal.dataobject.sms.sms_channel_do import SmsChannelDO
    from module_system.dal.dataobject.sms.sms_template_do import SmsTemplateDO
    from module_system.dal.mapper.sms.sms_channel_mapper import SmsChannelMapper
    from module_system.dal.mapper.sms.sms_template_mapper import SmsTemplateMapper
    from module_system.definitions.enums.sms.sms_send_status_enum import SmsSendStatusEnum

    code = "contract_" + uuid4().hex[:10]
    application = system_app.state.application_context
    with application.execution(), system_app.state.database.scope():
        security = application.container.get(SecurityService)
        async with security.authorized(
            admin_client.headers["Authorization"].removeprefix("Bearer "),
            RoutePolicy(tenant_required=True, realm=SecurityRealm.TENANT),
        ):
            channel = await application.container.get(SmsChannelMapper).insert(
                SmsChannelDO(
                    signature="Contract",
                    code="ALIYUN",
                    status=0,
                    api_key="test-only",
                    api_secret="test-only",
                )
            )
            await application.container.get(SmsTemplateMapper).insert(
                SmsTemplateDO(
                    type=1,
                    status=0,
                    code=code,
                    name="Contract",
                    content="Code {code}",
                    params=["code"],
                    api_template_id="test",
                    channel_id=channel.id,
                    channel_code="ALIYUN",
                )
            )
    result = (
        await admin_client.post(
            "/admin-api/system/sms/template/send-sms",
            json={
                "mobile": "13800138000",
                "templateCode": code,
                "templateParams": {"code": "1234"},
            },
        )
    ).json()
    assert result["code"] == 0, result
    with system_database[2].cursor() as cursor:
        cursor.execute(
            "SELECT template_content,send_status,user_id FROM system_sms_log WHERE id=%s",
            (int(result["data"]),),
        )
        assert cursor.fetchone() == ("Code 1234", SmsSendStatusEnum.IGNORE.code, None)


async def test_permission_cache_invalidation_obeys_transaction_commit(admin_client, system_app):
    from framework.starter_cache.core.cache_handler import CacheHandler
    from module_system.dal.cache.cache_key_constants import SystemCacheKeys
    from module_system.service.permission.permission_cache_service import PermissionCacheService

    application = system_app.state.application_context
    with application.execution(), system_app.state.database.scope():
        security = application.container.get(SecurityService)
        async with security.authorized(
            admin_client.headers["Authorization"].removeprefix("Bearer "),
            RoutePolicy(tenant_required=True, realm=SecurityRealm.TENANT),
        ):
            cache = application.container.get(CacheHandler)
            service = application.container.get(PermissionCacheService)
            key = "contract-" + uuid4().hex
            await cache.set(SystemCacheKeys.ROLE, key, None)
            with pytest.raises(RuntimeError, match="rollback"):
                async with system_app.state.database.transaction():
                    await service.invalidate_role_caches()
                    assert (await cache.get(SystemCacheKeys.ROLE, key)).hit
                    raise RuntimeError("rollback")
            assert (await cache.get(SystemCacheKeys.ROLE, key)).hit
            async with system_app.state.database.transaction():
                await service.invalidate_role_caches()
            assert not (await cache.get(SystemCacheKeys.ROLE, key)).hit


async def test_manual_announcement_publish_is_atomic_and_does_not_mutate_template(
    admin_client, system_app, monkeypatch
):
    from framework.common.exception.exceptions.service_exception import ServiceException
    from module_system.dal.dataobject.announcement.announcement_do import AnnouncementDO
    from module_system.dal.dataobject.notification.notice_log_do import NoticeLogDO
    from module_system.dal.dataobject.notification.notice_message_do import NoticeMessageDO
    from module_system.dal.mapper.announcement.announcement_mapper import AnnouncementMapper
    from module_system.dal.mapper.notification.notice_log_mapper import NoticeLogMapper
    from module_system.dal.mapper.notification.notice_mapper import NoticeMapper
    from module_system.dal.mapper.notification.notice_message_mapper import NoticeMessageMapper
    from module_system.service.announcement.announcement_service import AnnouncementService

    application = system_app.state.application_context
    with application.execution(), system_app.state.database.scope():
        security = application.container.get(SecurityService)
        async with security.authorized(
            admin_client.headers["Authorization"].removeprefix("Bearer "),
            RoutePolicy(tenant_required=True, realm=SecurityRealm.TENANT),
        ):
            mapper = application.container.get(AnnouncementMapper)
            service = application.container.get(AnnouncementService)
            notices = application.container.get(NoticeMapper)
            messages = application.container.get(NoticeMessageMapper)
            logs = application.container.get(NoticeLogMapper)
            template = await notices.select_by_code("system_announcement_publish")
            original_template = (template.title, template.content, list(template.channels))
            title = uuid4().hex + "公告" * 34
            announcement = await mapper.insert(
                AnnouncementDO(
                    title=title, content="Manual", status=0, category=1, publisher="test"
                )
            )
            assert await service.publish_announcement(announcement.id)
            count = await messages.count(NoticeMessageDO.notice_title == title)
            assert count > 0
            assert await logs.count(NoticeLogDO.notice_title == title) == 1
            assert not await service.publish_announcement(announcement.id)
            assert await messages.count(NoticeMessageDO.notice_title == title) == count
            template = await notices.select_by_code("system_announcement_publish")
            assert (template.title, template.content, template.channels) == original_template

            failed_title = "Rollback" + uuid4().hex
            failed = await mapper.insert(
                AnnouncementDO(
                    title=failed_title, content="Fail", status=0, category=1, publisher="test"
                )
            )
            original_send = service.notice_service.send_notice_direct

            async def fail_after_persistence(req):
                await original_send(req)
                raise RuntimeError("notification failure")

            monkeypatch.setattr(
                service.notice_service, "send_notice_direct", fail_after_persistence
            )
            with pytest.raises(RuntimeError, match="notification failure"):
                await service.publish_announcement(failed.id)
            assert (await mapper.select_by_id(failed.id)).status == 0
            assert await messages.count(NoticeMessageDO.notice_title == failed_title) == 0
            assert await logs.count(NoticeLogDO.notice_title == failed_title) == 0

            async def reject_dispatch(*args, **kwargs):
                return {int(security.context.require().account_id)}

            monkeypatch.setattr(service.notice_service, "send_notice_direct", original_send)
            monkeypatch.setattr(
                service.notice_service.notification_dispatcher, "send_notification", reject_dispatch
            )
            with pytest.raises(ServiceException):
                await service.publish_announcement(failed.id)
            assert (await mapper.select_by_id(failed.id)).status == 0
            assert await logs.count(NoticeLogDO.notice_title == failed_title) == 0


async def test_notice_message_api_marks_only_the_requested_user_messages(admin_client, system_app):
    from module_system.api.notification.dto.notice_message_mark_read_req_dto import (
        NoticeMessageMarkReadReqDTO,
    )
    from module_system.api.notification.notice_message_api import NoticeMessageApi
    from module_system.dal.mapper.notification.notice_mapper import NoticeMapper
    from module_system.service.notification.notice_message_service import NoticeMessageService

    application = system_app.state.application_context
    with application.execution(), system_app.state.database.scope():
        security = application.container.get(SecurityService)
        async with security.authorized(
            admin_client.headers["Authorization"].removeprefix("Bearer "),
            RoutePolicy(tenant_required=True, realm=SecurityRealm.TENANT),
        ) as identity:
            user_id = int(identity.account_id)
            notice = await application.container.get(NoticeMapper).select_by_code(
                "system_announcement_publish"
            )
            message_id = await application.container.get(
                NoticeMessageService
            ).create_notice_message(
                NoticeMessageCreateBO(
                    user_id=user_id,
                    user_type=2,
                    notice=notice,
                    sent_channels=["INTERNAL"],
                    publisher_info=None,
                )
            )
            api = application.container.get(NoticeMessageApi)
            assert (
                await api.mark_read(
                    NoticeMessageMarkReadReqDTO(
                        user_id=user_id + 1, user_type=2, notice_message_id=message_id
                    )
                )
                == 0
            )
            request = NoticeMessageMarkReadReqDTO(
                user_id=user_id, user_type=2, notice_message_id=message_id
            )
            assert await api.mark_read(request) == 1
            assert await api.mark_read(request) == 0


@pytest.mark.parametrize(
    "path",
    [
        "user/page",
        "user/profile/get",
        "dept/list",
        "dept/post/page",
        "permission/role/page",
        "permission/menu/list",
        "dict/type/page",
        "dict/data/page",
        "tenant/page",
        "tenant/package/page",
        "logger/login-log/page",
        "logger/operate-log/page",
        "mail/account/page",
        "mail/template/page",
        "mail/log/page",
        "sms/channel/page",
        "sms/template/page",
        "sms/log/page",
        "social/client/page",
        "social/user/page",
        "oauth2/client/page",
        "oauth2/token/page",
        "notification/page",
        "announcement/page",
    ],
)
async def test_read_endpoints(admin_client, path):
    response = await admin_client.get("/admin-api/system/" + path)
    assert response.status_code == 200, response.text
    assert response.json()["code"] == 0, response.json()


async def test_department_and_user_crud(admin_client):
    suffix = uuid4().hex[:10]
    department = {"name": "Test" + suffix, "parentId": "0", "sort": 1, "status": 1}
    result = (await admin_client.post("/admin-api/system/dept/create", json=department)).json()
    assert result["code"] == 0, result
    identifier = result["data"]
    assert isinstance(identifier, str)
    user = {
        "username": "u" + suffix,
        "nickname": "Tester",
        "password": "Password123",
        "deptId": identifier,
        "postIds": [],
    }
    result = (await admin_client.post("/admin-api/system/user/create", json=user)).json()
    assert result["code"] == 0, result
    user_id = result["data"]
    assert isinstance(user_id, str)
    found = (await admin_client.get("/admin-api/system/user/get", params={"id": user_id})).json()
    assert found["code"] == 0, found
    assert found["data"]["deptId"] == identifier
    result = (
        await admin_client.delete("/admin-api/system/user/delete", params={"id": user_id})
    ).json()
    assert result["code"] == 0, result
    result = (
        await admin_client.delete("/admin-api/system/dept/delete", params={"id": identifier})
    ).json()
    assert result["code"] == 0, result


async def test_rotation_replay_and_logout(system_app):
    async with AsyncClient(
        transport=ASGITransport(app=system_app), base_url="http://testserver"
    ) as client:
        login = (
            await client.post(
                "/admin-api/system/auth/login", json={"username": "admin", "password": "admin123"}
            )
        ).json()
        assert login["code"] == 0, login
        old_access = login["data"]["accessToken"]
        old_refresh = client.cookies.get("system_refresh")
        response = await client.post(
            "/admin-api/system/auth/refresh-token", headers={"Origin": "http://testserver"}
        )
        refreshed = response.json()
        assert refreshed["code"] == 0, refreshed
        assert client.cookies.get("system_refresh") != old_refresh
        old_result = (
            await client.get(
                "/admin-api/system/auth/codes", headers={"Authorization": "Bearer " + old_access}
            )
        ).json()
        assert old_result["code"] == SecurityErrorCodes.REVOKED.code
        new_access = refreshed["data"]["accessToken"]
        replay = (
            await client.post(
                "/admin-api/system/auth/refresh-token",
                headers={"Origin": "http://testserver", "Cookie": "system_refresh=" + old_refresh},
            )
        ).json()
        assert replay["code"] == SecurityErrorCodes.REVOKED.code, replay
        result = (
            await client.get(
                "/admin-api/system/auth/codes", headers={"Authorization": "Bearer " + new_access}
            )
        ).json()
        assert result["code"] == SecurityErrorCodes.REVOKED.code, result


async def test_openapi_and_public_boundaries(system_app):
    document = system_app.openapi()
    assert "/admin-api/system/auth/login" in document["paths"]
    assert not any(path.startswith("/weapp-api/system/") for path in document["paths"])
    async with AsyncClient(
        transport=ASGITransport(app=system_app), base_url="http://testserver"
    ) as client:
        response = await client.get("/admin-api/system/user/page")
        assert response.json()["code"] == SecurityErrorCodes.MISSING.code
        response = await client.post("/admin-api/system/auth/refresh-token")
        assert response.json()["code"] == SecurityErrorCodes.ORIGIN.code


async def test_role_assignment_and_permission_revocation(admin_client, system_app):
    suffix = uuid4().hex[:8]
    result = (
        await admin_client.post(
            "/admin-api/system/user/create",
            json={
                "username": "p" + suffix,
                "nickname": "Permission",
                "password": "Password123",
                "postIds": [],
            },
        )
    ).json()
    assert result["code"] == 0, result
    user_id = result["data"]
    result = (
        await admin_client.post(
            "/admin-api/system/permission/role/create",
            json={"name": "Role" + suffix, "code": "test" + suffix, "sort": 1},
        )
    ).json()
    assert result["code"] == 0, result
    role_id = result["data"]
    result = (
        await admin_client.post(
            "/admin-api/system/permission/assign-user-role",
            json={"userId": user_id, "roleIds": [role_id]},
        )
    ).json()
    assert result["code"] == 0, result
    async with AsyncClient(
        transport=ASGITransport(app=system_app, client=("127.0.0.2", 10000)),
        base_url="http://testserver",
    ) as user:
        result = (
            await user.post(
                "/admin-api/system/auth/login",
                json={"username": "p" + suffix, "password": "Password123"},
            )
        ).json()
        assert result["code"] == 0, result
        user.headers["Authorization"] = "Bearer " + result["data"]["accessToken"]
        before = (await user.get("/admin-api/system/auth/get-permission-info")).json()
        assert before["code"] == 0, before
        assert "test" + suffix in before["data"]["roles"]
        assert (await user.get("/admin-api/system/user/page")).json()[
            "code"
        ] == SecurityErrorCodes.DENIED.code
        result = (
            await admin_client.post(
                "/admin-api/system/permission/assign-user-role",
                json={"userId": user_id, "roleIds": []},
            )
        ).json()
        assert result["code"] == 0, result
        after = (await user.get("/admin-api/system/auth/get-permission-info")).json()
        assert after["code"] == 0, after
        assert after["data"]["roles"] == []
        result = (
            await admin_client.put(
                "/admin-api/system/user/update-password",
                json={"id": user_id, "password": "NewPassword123"},
            )
        ).json()
        assert result["code"] == 0, result
        assert (await user.get("/admin-api/system/auth/codes")).json()[
            "code"
        ] == SecurityErrorCodes.CREDENTIALS.code


async def test_database_rollback_and_authentication_read_boundary(system_app, admin_client):
    from sqlalchemy import select, text

    from framework.starter_database.query.authentication_reader import AuthenticationReader
    from framework.starter_tenant.definitions.constants.tenant_error_codes import TenantErrorCodes
    from framework.starter_tenant.exception.tenant_exception import TenantException
    from module_system.dal.dataobject.dept.dept_do import DeptDO
    from module_system.dal.dataobject.user.admin_user_do import AdminUserDO
    from module_system.dal.mapper.dept.dept_mapper import DeptMapper

    application = system_app.state.application_context
    database = system_app.state.database
    with application.execution(), database.scope():
        security = application.container.get(SecurityService)
        async with security.authorized(
            admin_client.headers["Authorization"].removeprefix("Bearer "),
            RoutePolicy(tenant_required=True, realm=SecurityRealm.TENANT),
        ):
            mapper = application.container.get(DeptMapper)
            name = "rollback" + uuid4().hex[:8]
            with pytest.raises(RuntimeError, match="rollback-marker"):
                async with database.transaction():
                    entry = await mapper.insert(DeptDO(name=name, parent_id=0, sort=0, status=1))
                    assert entry.id and entry.tenant_id == "1"
                    assert entry.creator == security.context.require().account_id
                    raise RuntimeError("rollback-marker")
            assert await mapper.select_list(DeptDO.name == name) == []
            with pytest.raises(TenantException) as cross_tenant:
                await mapper.insert(
                    DeptDO(name=name, parent_id=0, sort=0, status=1, tenant_id="999")
                )
            assert cross_tenant.value.error_code is TenantErrorCodes.WRITE
        reader = AuthenticationReader(database, (AdminUserDO,), tenant_id="1")
        with pytest.raises(ValueError):
            await reader.read(select(DeptDO.__table__))
        with pytest.raises(ValueError):
            await reader.read(text("SELECT * FROM system_users"))
        with pytest.raises(TenantException) as missing_tenant:
            async with database.read_session() as session:
                await session.execute(select(AdminUserDO))
        assert missing_tenant.value.error_code is TenantErrorCodes.MISSING


async def test_excel_export_and_snowflake_input(admin_client):
    response = await admin_client.get("/admin-api/system/user/export-excel")
    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"], response.text
    assert response.content.startswith(b"PK")
    response = await admin_client.put(
        "/admin-api/system/dept/update-status", json={"id": 10100000020001, "status": 1}
    )
    assert response.json()["code"] == 422


async def test_audit_records_are_persistent_and_hide_passwords(system_database):
    _, _, connection = system_database
    with connection.cursor() as cursor:
        cursor.execute(
            "SELECT result, action, extra FROM system_operate_log WHERE result = 'success'"
        )
        rows = cursor.fetchall()
    assert rows
    assert all(
        "Password123" not in action + extra and "$2b$" not in action + extra
        for _, action, extra in rows
    )


async def test_oauth2_password_scope_and_code_consumption(admin_client, system_app):
    import base64
    from urllib.parse import parse_qs, urlsplit

    credential = base64.b64encode(b"default:dushan-admin-secret").decode()
    async with AsyncClient(
        transport=ASGITransport(app=system_app),
        base_url="http://testserver",
        headers={"Authorization": "Basic " + credential},
    ) as client:
        request = {
            "grant_type": "password",
            "username": "admin",
            "password": "admin123",
            "scope": "user.read",
        }
        result = (await client.post("/admin-api/system/oauth2/open/token", data=request)).json()
        assert result["code"] == 0, result
        assert result["data"]["scope"] == "user.read"
        rejected = (
            await client.post(
                "/admin-api/system/oauth2/open/token", data={**request, "scope": "not.granted"}
            )
        ).json()
        assert rejected["code"] != 0
        authorization = (
            await admin_client.post(
                "/admin-api/system/oauth2/open/authorize",
                params={
                    "client_id": "default",
                    "redirect_uri": "http://localhost",
                    "auto_approve": "true",
                    "response_type": "code",
                    "scope": '{"user.read":true}',
                    "state": "test-state",
                },
            )
        ).json()
        assert authorization["code"] == 0, authorization
        code = parse_qs(urlsplit(authorization["data"]).query)["code"][0]
        exchange = {
            "grant_type": "authorization_code",
            "code": code,
            "redirect_uri": "http://localhost",
            "state": "test-state",
        }
        result = (await client.post("/admin-api/system/oauth2/open/token", data=exchange)).json()
        assert result["code"] == 0, result
        assert (await client.post("/admin-api/system/oauth2/open/token", data=exchange)).json()[
            "code"
        ] != 0


async def test_socket_ticket_is_single_use_and_logout_revokes_family(admin_client, system_app):
    from framework.starter_security.exception.security_exception import SecurityException
    from module_system.service.oauth2.oauth2_token_service import OAuth2TokenService

    response = (await admin_client.post("/admin-api/system/auth/websocket-ticket")).json()
    assert response["code"] == 0, response
    application = system_app.state.application_context
    with application.execution(), system_app.state.database.scope():
        service = application.container.get(OAuth2TokenService)
        security = application.container.get(SecurityService)
        identity = await service.consume_socket_ticket(
            response["data"], application_id=security.settings.application_id, domain="admin"
        )
        assert identity.tenant_id == "1"
        with pytest.raises(SecurityException):
            await service.consume_socket_ticket(
                response["data"], application_id=security.settings.application_id, domain="admin"
            )


async def test_live_http_process(system_app):
    import asyncio
    import os
    import socket
    import subprocess
    import sys
    from pathlib import Path

    folder = Path(system_app.state.bootstrap.base_dir)
    temp = folder / "Temp/live-http"
    temp.mkdir(parents=True)
    import yaml

    values = yaml.safe_load((folder / "application.yaml").read_text(encoding="utf-8"))
    values["config"]["models"]["database"]["snowflake_machine_id"] = 914
    (temp / "application.yaml").write_text(
        yaml.safe_dump(values, allow_unicode=True), encoding="utf-8"
    )
    with socket.socket() as listener:
        listener.bind(("127.0.0.1", 0))
        port = listener.getsockname()[1]
    environment = dict(os.environ)
    environment.update(
        TEMP=str(temp),
        TMP=str(temp),
        TMPDIR=str(temp),
        DUSHAN_CONFIG_DIR=str(temp),
        PYTHONDONTWRITEBYTECODE="1",
        PYTHONPATH=str(Path(__file__).resolve().parents[2] / "dushan-admin-backend"),
    )
    with (temp / "server.log").open("wb") as output:
        process = subprocess.Popen(
            [
                sys.executable,
                "-B",
                "-m",
                "uvicorn",
                "server.starter_server:app",
                "--host",
                "127.0.0.1",
                "--port",
                str(port),
                "--log-level",
                "error",
            ],
            cwd=folder,
            env=environment,
            stdout=output,
            stderr=subprocess.STDOUT,
            creationflags=subprocess.CREATE_NO_WINDOW,
        )
        try:
            async with AsyncClient(base_url=f"http://127.0.0.1:{port}", timeout=5) as client:
                for _ in range(200):
                    assert process.poll() is None, (temp / "server.log").read_text(encoding="utf-8")
                    try:
                        response = await client.get("/health")
                        if response.status_code == 200:
                            break
                    except Exception as error:
                        from httpx import ConnectError

                        if not isinstance(error, ConnectError):
                            raise
                    await asyncio.sleep(0.1)
                else:
                    pytest.fail("live server did not become ready")
                response = await client.post(
                    "/admin-api/system/auth/login",
                    json={"username": "admin", "password": "admin123"},
                )
                assert response.json()["code"] == 0, response.text
                response = await client.get(
                    "/admin-api/system/user/profile/get",
                    headers={"Authorization": "Bearer " + response.json()["data"]["accessToken"]},
                )
                assert response.json()["code"] == 0, response.text
        finally:
            process.terminate()
            process.wait(timeout=20)


async def test_tenant_provision_and_cross_tenant_isolation(admin_client, system_app):
    from framework.starter_cache.core.cache_handler import CacheHandler
    from module_system.controller.admin.tenant.vo.tenant.tenant_save_req_vo import TenantSaveReqVO
    from module_system.dal.cache.cache_key_constants import SystemCacheKeys
    from module_system.dal.mapper.user.admin_user_mapper import AdminUserMapper
    from module_system.service.auth.system_workload_service import SystemWorkloadService
    from module_system.service.tenant.tenant_service import TenantService

    suffix = uuid4().hex[:8]
    package = (
        await admin_client.post(
            "/admin-api/system/tenant/package/create",
            json={"name": "Package" + suffix, "status": 1, "menuIds": []},
        )
    ).json()
    assert package["code"] == 0, package
    application = system_app.state.application_context
    with application.execution(), system_app.state.database.scope():
        security = application.container.get(SecurityService)
        async with security.authorized(
            admin_client.headers["Authorization"].removeprefix("Bearer "),
            RoutePolicy(tenant_required=True, realm=SecurityRealm.TENANT),
        ):
            identifier = await application.container.get(TenantService).create_tenant(
                TenantSaveReqVO.model_validate(
                    {
                        "name": "Tenant" + suffix,
                        "contactName": "Owner",
                        "status": 1,
                        "websites": [],
                        "packageId": package["data"],
                        "expireTime": "2099-01-01T00:00:00",
                        "accountCount": 10,
                        "username": "admin",
                        "password": "Password123",
                    }
                )
            )
        tenant = {"data": str(identifier)}
        workloads = application.container.get(SystemWorkloadService)
        cache = application.container.get(CacheHandler)
        async with workloads.scope("system.auth", tenant["data"]):
            await cache.set(SystemCacheKeys.ROLE, "missing-role", None)
            cached = await cache.get(SystemCacheKeys.ROLE, "missing-role")
            assert cached.hit and cached.value is None
            redis = cache.get_client(SystemCacheKeys.ROLE)
            ttl = await redis.ttl(cache.build_full_key(SystemCacheKeys.ROLE, "missing-role"))
            assert 0 < ttl <= 3600
            users = await application.container.get(AdminUserMapper).select_list()
            assert len(users) == 1
            assert users[0].tenant_id == tenant["data"]
            foreign_user = str(users[0].id)
        async with workloads.scope("system.auth", "1"):
            assert not (await cache.get(SystemCacheKeys.ROLE, "missing-role")).hit
            assert (
                await application.container.get(AdminUserMapper).select_by_id(int(foreign_user))
                is None
            )
    response = (
        await admin_client.get("/admin-api/system/user/get", params={"id": foreign_user})
    ).json()
    assert response["code"] == 0 and response["data"] is None, response


async def test_message_proofs_bind_payload_audience_and_workload(admin_client, system_app):
    from framework.starter_security.core.opaque_token import OpaqueToken
    from framework.starter_security.exception.security_exception import SecurityException
    from module_system.service.auth.system_message_service import SystemMessageService
    from module_system.service.auth.system_workload_service import SystemWorkloadService
    from module_system.service.oauth2.oauth2_token_service import OAuth2TokenService

    application = system_app.state.application_context
    with application.execution(), system_app.state.database.scope():
        messages = application.container.get(SystemMessageService)
        security = application.container.get(SecurityService)
        metadata = {"application_id": security.settings.application_id, "domain": "admin"}
        token = admin_client.headers["Authorization"].removeprefix("Bearer ")
        session = await application.container.get(OAuth2TokenService).resolve_session(
            OpaqueToken.digest(token), **metadata
        )
        proof = await messages.issue(session, b"mail-record-id", audience="system.mail.send")
        assert (
            await messages.verify(proof, b"mail-record-id", audience="system.mail.send", **metadata)
        ).family_id == session.family_id
        with pytest.raises(SecurityException):
            await messages.verify(proof, b"changed-record", audience="system.mail.send", **metadata)
        with pytest.raises(SecurityException):
            await messages.verify(proof, b"mail-record-id", audience="system.sms.send", **metadata)
        workload = await application.container.get(SystemWorkloadService).authenticate(
            "module_system", capability="system.mail.send", tenant_id="1", **metadata
        )
        proof = await messages.issue_workload(
            workload, b"mail-record-id", audience="system.mail.send", capability="system.mail.send"
        )
        received = await messages.verify_workload(
            proof, b"mail-record-id", audience="system.mail.send", **metadata
        )
        assert received.identity.tenant_id == "1"
        assert received.capability == "system.mail.send"
        with pytest.raises(SecurityException):
            await messages.issue_workload(
                workload,
                b"mail-record-id",
                audience="system.mail.send",
                capability="infra.arbitrary",
            )
