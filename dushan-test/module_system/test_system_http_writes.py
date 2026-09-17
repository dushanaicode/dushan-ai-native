import re
from uuid import uuid4

import pytest
from fastapi.routing import APIRoute, _iter_routes_with_context
from httpx import ASGITransport, AsyncClient

from framework.starter_security.core.security_service import SecurityService
from framework.starter_security.enums.security_realm import SecurityRealm
from framework.starter_web.routing.route_policy import RoutePolicy
from module_system.dal.mapper.notification.notice_mapper import NoticeMapper
from module_system.framework.sms.client.providers.aliyun_sms_client import AliyunSmsClient
from module_system.framework.sms.model.sms_template_resp_dto import SmsTemplateRespDTO
from module_system.router import routers
from module_system.service.notification.bo.notice_message_create_bo import NoticeMessageCreateBO
from module_system.service.notification.notice_message_service import NoticeMessageService

pytestmark = pytest.mark.asyncio(loop_scope="module")

CASES = [
    ("announcement", "system_announcement"),
    ("dept", "system_dept"),
    ("dept/post", "system_post"),
    ("dict/type", "system_dict_type"),
    ("dict/data", "system_dict_data"),
    ("mail/account", "system_mail_account"),
    ("mail/template", "system_mail_template"),
    ("notification", "system_notification_notice"),
    ("oauth2/client", "system_oauth2_client"),
    ("permission/menu", "system_menu"),
    ("permission/role", "system_role"),
    ("sms/channel", "system_sms_channel"),
    ("sms/template", "system_sms_template"),
    ("social/client", "system_social_client"),
    ("tenant", "system_tenant"),
    ("tenant/package", "system_tenant_package"),
    ("user", "system_users"),
]


async def create(client, domain, payload):
    result = (await client.post(f"/admin-api/system/{domain}/create", json=payload)).json()
    assert result["code"] == 0, (domain, result)
    assert isinstance(result["data"], str)
    return result["data"]


async def payload_for(client, domain):
    key = "http_" + uuid4().hex[:12]
    match domain:
        case "announcement":
            return {
                "title": key,
                "content": "Test",
                "status": 0,
                "category": 1,
                "publisher": "Tester",
            }
        case "dept":
            return {"name": key, "parentId": "0", "sort": 1, "status": 1}
        case "dept/post":
            return {"name": key, "code": key, "sort": 1, "status": 1}
        case "dict/type":
            return {"name": key, "type": key, "status": 1}
        case "dict/data":
            await create(client, "dict/type", {"name": key, "type": key, "status": 1})
            return {"label": key, "value": key, "dictType": key, "sort": 1, "status": 1}
        case "mail/account":
            return {
                "mail": key + "@example.test",
                "username": key,
                "password": "test-only",
                "host": "smtp.example.test",
                "port": 465,
                "sslEnable": True,
                "starttlsEnable": False,
            }
        case "mail/template":
            account = await create(
                client, "mail/account", await payload_for(client, "mail/account")
            )
            return {
                "name": key,
                "code": key,
                "accountId": account,
                "nickname": "Tester",
                "title": "Hello",
                "content": "Hello {name}",
                "status": 0,
            }
        case "notification":
            return {
                "title": key,
                "content": "Test",
                "type": 2,
                "userType": 2,
                "channels": ["INTERNAL"],
                "publisher": "Tester",
                "status": 1,
            }
        case "oauth2/client":
            return {
                "clientId": key,
                "name": key,
                "secret": "test-only",
                "logo": "https://example.test/logo.png",
                "status": 1,
                "userType": 2,
                "accessTokenValiditySeconds": 3600,
                "refreshTokenValiditySeconds": 7200,
                "redirectUris": ["https://example.test/callback"],
                "authorizedGrantTypes": ["password"],
                "scopes": [],
                "autoApproveScopes": [],
                "authorities": [],
                "resourceIds": [],
            }
        case "permission/menu":
            return {
                "name": key,
                "kind": "action",
                "permission": "test:" + key,
                "parentId": "0",
                "sort": 1,
                "status": 1,
            }
        case "permission/role":
            return {"name": key, "code": key, "sort": 1}
        case "sms/channel":
            return {
                "signature": "t" + key[-8:],
                "code": "ALIYUN",
                "status": 1,
                "apiKey": "test-only",
                "apiSecret": "test-only",
            }
        case "sms/template":
            channel = await create(client, "sms/channel", await payload_for(client, "sms/channel"))
            return {
                "name": key,
                "code": key,
                "type": 1,
                "status": 0,
                "content": "Code {code}",
                "apiTemplateId": "test-template",
                "channelId": channel,
            }
        case "social/client":
            return {
                "name": key,
                "socialType": 31,
                "userType": 2,
                "clientId": key,
                "clientSecret": "test-only",
                "status": 0,
            }
        case "tenant":
            package = await create(
                client, "tenant/package", await payload_for(client, "tenant/package")
            )
            return {
                "name": key,
                "contactName": "Tester",
                "status": 1,
                "websites": [],
                "packageId": package,
                "expireTime": "2099-01-01T00:00:00",
                "accountCount": 10,
                "username": "admin",
                "password": "Password123",
            }
        case "tenant/package":
            return {"name": key, "status": 1, "menuIds": []}
        case "user":
            return {
                "username": "u" + uuid4().hex[:10],
                "nickname": key,
                "password": "Password123",
                "postIds": [],
            }
    raise AssertionError(domain)


@pytest.mark.parametrize("domain,table", CASES)
async def test_crud_and_batch_delete_keep_unselected_rows(
    admin_client, system_app, system_database, monkeypatch, domain, table
):
    async def template(self, identifier):
        return SmsTemplateRespDTO(id=identifier, content="Code {code}", audit_status=2)

    monkeypatch.setattr(AliyunSmsClient, "get_sms_template", template)
    payloads = [await payload_for(admin_client, domain) for _ in range(3)]
    if domain == "social/client":
        with system_database[2].cursor() as cursor:
            cursor.execute(
                "SELECT social_type FROM system_social_client WHERE user_type=2 AND tenant_id='1' AND deleted=0"
            )
            used = {row[0] for row in cursor.fetchall()}
        providers = [value for value in (10, 20, 31, 32, 33) if value not in used]
        assert len(providers) >= 3
        for payload, provider in zip(payloads, providers[:3], strict=True):
            payload["socialType"] = provider
    ids = [await create(admin_client, domain, payload) for payload in payloads]
    path = "/admin-api/system/" + domain
    update = {**payloads[0], "id": ids[0]}
    if domain == "tenant":
        del update["username"], update["password"]
    field = next(
        name
        for name in ("nickname", "name", "title", "label", "signature", "username")
        if name in update
    )
    update[field] = "upd" + uuid4().hex[:8]
    changed = (await admin_client.put(path + "/update", json=update)).json()
    assert changed["code"] == 0, (domain, changed)
    found = (await admin_client.get(path + "/get", params={"id": ids[0]})).json()
    assert found["code"] == 0 and found["data"][field] == update[field], (domain, found)
    if domain not in {"announcement", "mail/account"}:
        status = 0 if payloads[0].get("status", 1) == 1 else 1
        changed_status = (
            await admin_client.put(path + "/update-status", json={"id": ids[0], "status": status})
        ).json()
        assert changed_status["code"] == 0, (domain, changed_status)
        found = (await admin_client.get(path + "/get", params={"id": ids[0]})).json()
        assert found["data"]["status"] == status, (domain, found)
    async with AsyncClient(
        transport=ASGITransport(app=system_app), base_url="http://testserver"
    ) as public:
        denied = (
            await public.delete(path + "/delete-list", params=[("ids", item) for item in ids[:2]])
        ).json()
        assert denied["code"] == 401, (domain, denied)
    for params in ({}, {"ids": ",".join(ids)}, [("ids", ids[0]), ("ids", "invalid")]):
        invalid = (await admin_client.delete(path + "/delete-list", params=params)).json()
        assert invalid["code"] == 422, (domain, invalid)
    with system_database[2].cursor() as cursor:
        cursor.execute(f"SELECT id FROM {table} WHERE id IN (%s,%s,%s) AND deleted=0", ids)
        assert {str(row[0]) for row in cursor.fetchall()} == set(ids)
    result = (
        await admin_client.delete(path + "/delete-list", params=[("ids", item) for item in ids[:2]])
    ).json()
    assert result["code"] == 0 and result["data"] == 2, (domain, result)
    with system_database[2].cursor() as cursor:
        cursor.execute(f"SELECT id FROM {table} WHERE id IN (%s,%s,%s) AND deleted=0", ids)
        assert [str(row[0]) for row in cursor.fetchall()] == ids[2:]
    single = (await admin_client.delete(path + "/delete", params={"id": ids[2]})).json()
    assert single["code"] == 0, (domain, single)


async def test_batch_token_revoke_only_selected_sessions(admin_client, system_app, system_database):
    payload = await payload_for(admin_client, "user")
    user_id = await create(admin_client, "user", payload)
    tokens = []
    async with AsyncClient(
        transport=ASGITransport(app=system_app), base_url="http://testserver"
    ) as user:
        for _ in range(3):
            login = (
                await user.post(
                    "/admin-api/system/auth/login",
                    json={"username": payload["username"], "password": payload["password"]},
                )
            ).json()
            assert login["code"] == 0, login
            tokens.append(login["data"]["accessToken"])
        with system_database[2].cursor() as cursor:
            cursor.execute(
                "SELECT id FROM system_oauth2_access_token WHERE user_id=%s AND deleted=0 ORDER BY id",
                (user_id,),
            )
            ids = [str(row[0]) for row in cursor.fetchall()]
        assert len(ids) == 3
        path = "/admin-api/system/oauth2/token/delete-list"
        denied = (await user.delete(path, params={"ids": ids[0]})).json()
        assert denied["code"] == 401, denied
        for params in ({}, {"ids": ",".join(ids)}, {"ids": "invalid"}):
            invalid = (await admin_client.delete(path, params=params)).json()
            assert invalid["code"] == 422, invalid
        result = (
            await admin_client.delete(path, params=[("ids", value) for value in ids[:2]])
        ).json()
        assert result["code"] == 0 and result["data"] == 2, result
        for token, expected in zip(tokens, (401, 401, 0), strict=True):
            response = (
                await user.get(
                    "/admin-api/system/auth/codes", headers={"Authorization": "Bearer " + token}
                )
            ).json()
            assert response["code"] == expected, response


async def test_http_mark_read_enforces_message_owner(admin_client, system_app, system_database):
    application = system_app.state.application_context
    with application.execution(), system_app.state.database.scope():
        async with application.container.get(SecurityService).authorized(
            admin_client.headers["Authorization"].removeprefix("Bearer "),
            RoutePolicy(tenant_required=True, realm=SecurityRealm.TENANT),
        ) as identity:
            notice = await application.container.get(NoticeMapper).select_by_code(
                "system_announcement_publish"
            )
            service = application.container.get(NoticeMessageService)
            user_id = int(identity.account_id)
            ids = [
                await service.create_notice_message(
                    NoticeMessageCreateBO(
                        user_id=owner,
                        user_type=2,
                        notice=notice,
                        sent_channels=["INTERNAL"],
                        publisher_info=None,
                    )
                )
                for owner in (user_id, user_id, user_id, user_id + 100000)
            ]
    response = (
        await admin_client.put(
            "/admin-api/system/notification/message/update-read",
            json={"ids": [str(value) for value in (ids[0], ids[1], ids[3])]},
        )
    ).json()
    assert response["code"] == 0, response
    with system_database[2].cursor() as cursor:
        cursor.execute(
            "SELECT id,read_status FROM system_notification_message WHERE id IN (%s,%s,%s,%s)", ids
        )
        actual = dict(cursor.fetchall())
    assert [bool(actual[value]) for value in ids] == [True, True, False, False]


@pytest.mark.parametrize(
    "domain,table", [("dept", "system_dept"), ("permission/menu", "system_menu")]
)
async def test_batch_delete_preserves_children_and_rolls_back(
    admin_client, system_database, domain, table
):
    from module_system.definitions.constants.error_code_constants import ErrorCodeConstants

    parent_payload = await payload_for(admin_client, domain)
    if domain == "permission/menu":
        parent_payload.update(kind="group", path="/" + uuid4().hex)
    parent = await create(admin_client, domain, parent_payload)
    child = await create(
        admin_client, domain, {**await payload_for(admin_client, domain), "parentId": parent}
    )
    unrelated = await create(admin_client, domain, await payload_for(admin_client, domain))
    path = "/admin-api/system/" + domain + "/delete-list"
    rejected = (
        await admin_client.delete(path, params=[("ids", unrelated), ("ids", parent)])
    ).json()
    error = (
        ErrorCodeConstants.DEPT_EXITS_CHILDREN
        if domain == "dept"
        else ErrorCodeConstants.MENU_EXISTS_CHILDREN
    )
    assert rejected["code"] == error.code, rejected
    with system_database[2].cursor() as cursor:
        cursor.execute(
            f"SELECT id FROM {table} WHERE id IN (%s,%s,%s) AND deleted=0",
            (parent, child, unrelated),
        )
        assert len(cursor.fetchall()) == 3
    removed = (await admin_client.delete(path, params=[("ids", parent), ("ids", child)])).json()
    assert removed["code"] == 0 and removed["data"] == 2, removed
    with system_database[2].cursor() as cursor:
        cursor.execute(
            f"SELECT id FROM {table} WHERE id IN (%s,%s,%s) AND deleted=0",
            (parent, child, unrelated),
        )
        assert [str(row[0]) for row in cursor.fetchall()] == [unrelated]


@pytest.mark.parametrize(
    "path",
    [
        "dept/post",
        "dict/data",
        "dict/type",
        "logger/login-log",
        "logger/operate-log",
        "mail/account",
        "mail/log",
        "mail/template",
        "permission/role",
        "sms/log",
        "sms/template",
        "tenant",
        "tenant/package",
        "user",
    ],
)
async def test_export_http_returns_readable_workbook(admin_client, path):
    from io import BytesIO

    from openpyxl import load_workbook

    response = await admin_client.get("/admin-api/system/" + path + "/export-excel")
    assert response.status_code == 200 and response.content.startswith(b"PK"), response.text[:200]
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    try:
        assert len(next(workbook.active.iter_rows(values_only=True))) > 0
    finally:
        workbook.close()


@pytest.mark.parametrize(
    "path",
    [
        "dept",
        "dept/post",
        "dict/data",
        "dict/type",
        "mail/account",
        "mail/template",
        "permission/menu",
        "permission/role",
        "sms/channel",
        "sms/template",
        "tenant",
        "tenant/package",
        "user",
    ],
)
async def test_selector_http_returns_list(admin_client, path):
    response = (await admin_client.get("/admin-api/system/" + path + "/simple-list")).json()
    assert response["code"] == 0 and isinstance(response["data"], list), (path, response)


async def test_all_protected_system_routes_reject_anonymous(system_app):
    checked = 0
    async with AsyncClient(
        transport=ASGITransport(app=system_app), base_url="http://testserver"
    ) as public:
        for router in routers:
            for route, context in _iter_routes_with_context(router.routes):
                if not isinstance(route, APIRoute):
                    continue
                policy = getattr(route.endpoint, RoutePolicy.ATTRIBUTE)
                if not policy.requires_identity:
                    continue
                path = context.path if context else route.path
                url = re.sub(r"\{[^}]+\}", "9223372036854775807", path)
                for method in route.methods:
                    response = (await public.request(method, url, json={})).json()
                    assert response["code"] == 401, (method, path, response)
                    checked += 1
    assert checked > 200
