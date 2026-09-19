import asyncio
import io
import struct
from datetime import date, datetime, timezone
from decimal import Decimal
from types import SimpleNamespace
from typing import Annotated
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.xml import functions as xml_functions
from pydantic import BaseModel, ConfigDict, Field, PlainSerializer

from fixtures.config_factory import ConfigFactory
from framework.common.enums.base_enum import BaseEnum
from framework.starter_excel.config.excel_settings import ExcelSettings
from framework.starter_excel.converter.area_converter import AreaConverter
from framework.starter_excel.converter.dict_converter import DictConverter
from framework.starter_excel.converter.enum_converter import EnumConverter
from framework.starter_excel.converter.ids_converter import IdsConverter
from framework.starter_excel.converter.json_converter import JsonConverter
from framework.starter_excel.converter.money_converter import MoneyConverter
from framework.starter_excel.core.excel_schema import ExcelSchema
from framework.starter_excel.definitions.constants.excel_error_codes import ExcelErrorCodes
from framework.starter_excel.exception.excel_exception import ExcelException
from framework.starter_excel.handler.excel_upload_validator import ExcelUploadValidator
from framework.starter_excel.model.conversion_context import ConversionContext
from framework.starter_excel.model.excel_column import ExcelColumn
from framework.starter_excel.model.excel_providers import ExcelProviders
from framework.starter_excel.reader import excel_reader as reader_module
from framework.starter_excel.reader.excel_reader import ExcelReader
from framework.starter_excel.service.dict_framework_service import DictFrameworkService
from framework.starter_excel.writer import excel_writer as writer_module
from framework.starter_excel.writer.excel_writer import ExcelWriter
from framework.starter_ip.model.area import Area
from framework.starter_ip.service.area_service import AreaService


def settings(**changes):
    values = ConfigFactory.values()["config"]["models"]["excel"]
    values.update(changes)
    return ExcelSettings.model_validate(values)


def upload(content, *, filename="data.xlsx", content_type=None):
    return SimpleNamespace(filename=filename, content_type=content_type, file=io.BytesIO(content))


def xlsx(rows):
    workbook = Workbook()
    try:
        for row in rows:
            workbook.active.append(row)
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()
    finally:
        workbook.close()


def replace_member(content, name, transform):
    output = io.BytesIO()
    with ZipFile(io.BytesIO(content)) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            target.writestr(entry.filename, transform(data) if entry.filename == name else data)
    return output.getvalue()


class State(BaseEnum):
    ENABLED = (1, "启用")
    DISABLED = (0, "禁用")


class Dictionaries:
    def __init__(self, label="中文"):
        self.calls = []
        self.label = label

    async def items(self, kind):
        self.calls.append(kind)
        return {"zh": self.label, "en": "English"}


class Names:
    def __init__(self):
        self.data = {1: "研发,一组", 2: '开发"二组'}
        self.name_calls = []
        self.id_calls = []

    async def names(self, ids):
        self.name_calls.append(list(ids))
        return {value: self.data[value] for value in ids if value in self.data}

    async def ids(self, names):
        self.id_calls.append(list(names))
        reverse = {name: value for value, name in self.data.items()}
        return {name: reverse[name] for name in names if name in reverse}


class TextRow(BaseModel):
    text: Annotated[str, ExcelColumn("文本")]


class FullRow(BaseModel):
    name: Annotated[str, ExcelColumn("名称")]
    id: Annotated[int, ExcelColumn("编号")]
    amount: Annotated[int, ExcelColumn("金额", converter=MoneyConverter())]
    details: Annotated[dict, ExcelColumn("JSON", converter=JsonConverter())]
    state: Annotated[int, ExcelColumn("状态", converter=EnumConverter(State))]
    language: Annotated[str, ExcelColumn("语言", converter=DictConverter("language"))]
    departments: Annotated[list[int], ExcelColumn("部门", converter=IdsConverter("departments"))]
    posts: Annotated[set[int], ExcelColumn("岗位", converter=IdsConverter("posts"))]
    area: Annotated[Area, ExcelColumn("地区", converter=AreaConverter())]
    day: Annotated[date, ExcelColumn("日期", number_format="yyyy-mm-dd")]
    instant: Annotated[datetime, ExcelColumn("时间")]
    precise: Annotated[Decimal, ExcelColumn("精确数值")]


@pytest.mark.asyncio
async def test_real_xlsx_roundtrip_all_converters_and_lookup_scope(tmp_path):
    areas = AreaService()
    areas.initialize()
    region = areas.get_area(110101)
    assert region is not None
    dictionaries, departments, posts = Dictionaries(), Names(), Names()
    providers = ExcelProviders(dictionaries, departments, posts, areas)
    row = FullRow(
        name="Alice",
        id=1234567890123456789,
        amount=-123,
        details={"内容": [1, False, None]},
        state=1,
        language="zh",
        departments=[2, 1],
        posts={1, 2},
        area=region,
        day=date(2026, 9, 12),
        instant=datetime(2026, 9, 12, 12, 30),
        precise=Decimal("1234567890123456.789"),
    )
    output = await ExcelWriter(settings()).write("数据", FullRow, [row, row], providers=providers)
    path = tmp_path / "roundtrip.xlsx"
    path.write_bytes(output.getvalue())
    with path.open("rb") as source:
        imported = await ExcelReader(settings()).read(
            SimpleNamespace(filename=path.name, content_type=None, file=source),
            FullRow,
            providers=providers,
        )
        assert not source.closed
        assert source.tell() == 0
    assert len(imported) == 2
    for actual in imported:
        assert actual.id == row.id and actual.amount == -123
        assert actual.details == row.details and actual.precise == row.precise
        assert actual.area is region
        assert actual.departments == [2, 1] and actual.posts == {1, 2}
        assert actual.day == row.day and actual.instant == row.instant
    assert dictionaries.calls == ["language", "language"]
    assert departments.name_calls == [[2, 1]] and departments.id_calls == [
        [posts.data[2], posts.data[1]]
    ]
    with ZipFile(path) as archive:
        assert b"1234567890123456789" in archive.read("xl/worksheets/sheet1.xml")
        assert b"<f>" not in archive.read("xl/worksheets/sheet1.xml")
    output.close()


@pytest.mark.asyncio
async def test_template_dropdowns_and_formula_text():
    values = ("=1+1", "+cmd", "-cmd", "@cmd", "\t=1", 'a,b"c')

    class OptionsRow(BaseModel):
        text: Annotated[str, ExcelColumn("文本", options=values)]

    writer = ExcelWriter(settings())
    output = await writer.write("_excel_options", OptionsRow, [OptionsRow(text=v) for v in values])
    workbook = load_workbook(output, data_only=False)
    try:
        assert [cell[0].value for cell in workbook.active.iter_rows(min_row=2)] == list(values)
        assert all(cell[0].data_type == "s" for cell in workbook.active.iter_rows(min_row=2))
        hidden = workbook.worksheets[1]
        assert hidden.sheet_state == "hidden"
        assert [row[0].value for row in hidden] == list(values)
        assert all(row[0].data_type == "s" for row in hidden)
        validation = next(iter(workbook.active.data_validations.dataValidation))
        assert validation.formula1 in workbook.defined_names
        assert str(validation.sqref) == "A2:A101"
    finally:
        workbook.close()
        output.close()
    template = await writer.template("模板", OptionsRow)
    assert await ExcelReader(settings()).read(upload(template.getvalue()), OptionsRow) == []
    template.close()


@pytest.mark.asyncio
async def test_formula_policy_error_coordinates_and_literal_mode():
    content = xlsx([["文本"], ["=1+1"]])
    with pytest.raises(ExcelException) as raised:
        await ExcelReader(settings()).read(upload(content), TextRow)
    assert [(issue.row, issue.column, issue.field) for issue in raised.value.issues] == [
        (2, 1, "text")
    ]
    assert isinstance(raised.value.__cause__, ValueError)
    actual = await ExcelReader(settings(formula_policy="literal")).read(upload(content), TextRow)
    assert actual[0].text == "=1+1"


@pytest.mark.asyncio
async def test_mapping_aliases_nulls_numbers_and_model_errors():
    class Row(BaseModel):
        model_config = ConfigDict(alias_generator=str.upper)
        mobile: Annotated[str, ExcelColumn("电话")]
        count: Annotated[int, ExcelColumn("数量")]
        enabled: Annotated[bool, ExcelColumn("布尔")]
        day: Annotated[date | None, ExcelColumn("日期")] = None

    content = xlsx([["布尔", "数量", "电话"], [False, 0, 123456], [None, None, None]])
    imported = await ExcelReader(settings()).read(upload(content), Row)
    assert len(imported) == 1
    assert imported[0].mobile == "123456" and imported[0].count == 0
    assert imported[0].enabled is False and imported[0].day is None
    content = xlsx([["布尔", "数量", "电话", "日期"], [True, True, "hello", 45000]])
    with pytest.raises(ExcelException) as raised:
        await ExcelReader(settings(max_errors=1)).read(upload(content), Row)
    assert len(raised.value.issues) == 1 and raised.value.issues[0].column == 2


@pytest.mark.asyncio
@pytest.mark.parametrize("headers", [["文本", "文本"], ["不存在"], []])
async def test_invalid_headers(headers):
    with pytest.raises(ExcelException) as raised:
        await ExcelReader(settings()).read(upload(xlsx([headers])), TextRow)
    assert raised.value.issues[0].row == 1


@pytest.mark.asyncio
async def test_unknown_column_policy_and_inherited_export_field_policy():
    class Parent(BaseModel):
        text: Annotated[str, ExcelColumn("文本")]
        password: Annotated[str, ExcelColumn("密码")]
        hidden: Annotated[str, ExcelColumn("隐藏"), Field(exclude=True)]

    class Child(Parent):
        amount: Annotated[int, ExcelColumn("金额", exportable=False)]

    assert ExcelWriter.export_fields(Parent) == [{"field": "text", "title": "文本"}]
    assert ExcelWriter.export_fields(Child) == ExcelWriter.export_fields(Parent)
    for fields in ([], ["password"], ["unknown"], ["text", "text"]):
        with pytest.raises(ExcelException):
            ExcelSchema(Child).export_columns(fields)
    imported = await ExcelReader(settings(unknown_columns="ignore")).read(
        upload(xlsx([["未知", "文本"], ["secret", "yes"]])),
        TextRow,
    )
    assert imported[0].text == "yes"


@pytest.mark.asyncio
async def test_field_serializers_are_preserved():
    class Row(BaseModel):
        text: Annotated[str, PlainSerializer(lambda _: "***"), ExcelColumn("文本")]

    result = await ExcelWriter(settings()).write("数据", Row, [Row(text="private")])
    assert (await ExcelReader(settings()).read(upload(result.getvalue()), TextRow))[0].text == "***"
    result.close()


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "value",
    [
        "0.001",
        "NaN",
        "Infinity",
        "oops",
        True,
        "1e999999",
        "1e-999999",
        "1e-999999999",
        "-1e-999999999",
    ],
)
async def test_money_rejects_lossy_or_invalid_values(value):
    context = ConversionContext(settings(), ExcelProviders())
    with pytest.raises((ValueError, ArithmeticError)):
        await MoneyConverter().to_python(value, context)


@pytest.mark.asyncio
async def test_money_extreme_zero_and_exact_cent_are_preserved():
    context = ConversionContext(settings(), ExcelProviders())
    assert await MoneyConverter().to_python("0e-999999999", context) == 0
    assert await MoneyConverter().to_python("-0e-999999999", context) == 0
    assert await MoneyConverter().to_python("0e999999999", context) == 0
    assert await MoneyConverter().to_python("-0e999999999", context) == 0
    assert await MoneyConverter().to_python("0.01", context) == 1


@pytest.mark.asyncio
@pytest.mark.parametrize("value", ['{"a":1,"a":2}', '{"a":NaN}', '{"a":1e999999}', "null", "bad"])
async def test_json_rejects_invalid_contract(value):
    with pytest.raises(ValueError):
        await JsonConverter().to_python(value, ConversionContext(settings(), ExcelProviders()))


@pytest.mark.asyncio
async def test_business_provider_missing_failure_and_duplicate_labels():
    class DictRow(BaseModel):
        value: Annotated[str, ExcelColumn("字典", converter=DictConverter("language"))]

    with pytest.raises(ExcelException) as raised:
        await ExcelWriter(settings()).write("数据", DictRow, [DictRow(value="zh")])
    assert raised.value.issues[0].field == "value"

    class Broken:
        async def items(self, kind):
            raise OSError("provider unavailable")

    with pytest.raises(ExcelException) as raised:
        await ExcelReader(settings()).read(
            upload(xlsx([["字典"], ["中文"]])),
            DictRow,
            providers=ExcelProviders(dictionaries=Broken()),
        )
    assert isinstance(raised.value.__cause__, OSError)
    assert raised.value.issues[0].column == 1

    class Duplicate:
        async def items(self, kind):
            return {"a": "same", "b": "same"}

    with pytest.raises(ValueError, match="重复"):
        await DictFrameworkService(Duplicate()).items("kind")
    service = DictFrameworkService(Dictionaries())
    assert await service.label("language", "unknown") is None
    assert await service.value("language", "中文") == "zh"


@pytest.mark.parametrize(
    "limit", ["max_upload_size_bytes", "max_archive_entries", "max_uncompressed_size_bytes"]
)
def test_upload_resource_limits(limit):
    with pytest.raises(ExcelException) as raised:
        ExcelUploadValidator(settings(**{limit: 1})).validate(upload(xlsx([["文本"], ["data"]])))
    assert raised.value.error_code == ExcelErrorCodes.LIMIT


@pytest.mark.parametrize("filename", ["book.xls", "book.csv", "book.xlsm", ""])
def test_only_real_xlsx_supported(filename):
    with pytest.raises(ExcelException):
        ExcelUploadValidator(settings()).validate(upload(xlsx([["文本"]]), filename=filename))


@pytest.mark.asyncio
async def test_untrusted_dimensions_and_sparse_far_coordinates():
    normal = xlsx([["文本"], ["included"]])
    content = replace_member(
        normal,
        "xl/worksheets/sheet1.xml",
        lambda data: data.replace(b'ref="A1:A2"', b'ref="A1:A1"'),
    )
    imported = await ExcelReader(settings()).read(upload(content), TextRow)
    assert imported[0].text == "included"
    for coordinate in (b"A900000", b"XFD2"):
        content = replace_member(
            normal,
            "xl/worksheets/sheet1.xml",
            lambda data: data.replace(b'r="A2"', b'r="' + coordinate + b'"'),
        )
        with pytest.raises(ExcelException) as raised:
            await ExcelReader(settings()).read(upload(content), TextRow)
        assert raised.value.error_code == ExcelErrorCodes.LIMIT
    content = replace_member(
        normal, "xl/worksheets/sheet1.xml", lambda data: data.replace(b' r="A2"', b"")
    )
    with pytest.raises(ExcelException) as raised:
        await ExcelReader(settings()).read(upload(content), TextRow)
    assert isinstance(raised.value.__cause__, KeyError)


def test_xml_dtd_and_forged_directory_are_rejected():
    assert xml_functions.DEFUSEDXML is True
    normal = xlsx([["文本"], ["data"]])
    content = replace_member(
        normal,
        "xl/workbook.xml",
        lambda data: b'<!DOCTYPE workbook [<!ENTITY bomb "explosion">]>' + data,
    )
    with pytest.raises(ExcelException) as raised:
        ExcelUploadValidator(settings()).validate(upload(content))
    assert raised.value.__cause__ is not None
    forged = bytearray(normal)
    eocd = forged.rfind(b"PK\x05\x06")
    struct.pack_into("<HH", forged, eocd + 8, 1, 1)
    with pytest.raises(ExcelException, match="XML 无效"):
        ExcelUploadValidator(settings()).validate(upload(forged))


@pytest.mark.asyncio
async def test_actual_iteration_and_text_limits():
    normal = xlsx([["文本"], ["1234"]])
    with pytest.raises(ExcelException):
        await ExcelReader(settings(max_cell_text_length=3)).read(upload(normal), TextRow)
    with pytest.raises(ExcelException):
        await ExcelWriter(settings(max_cells=1)).write("数据", TextRow, [TextRow(text="x")])
    with pytest.raises(ExcelException):
        await ExcelReader(settings(max_cells=1)).read(upload(normal), TextRow)


@pytest.mark.asyncio
async def test_reader_writer_cancellation_closes_owned_resources(monkeypatch):
    class DictRow(BaseModel):
        value: Annotated[str, ExcelColumn("字典", converter=DictConverter("language"))]

    started = asyncio.Event()

    class Blocking:
        async def items(self, kind):
            started.set()
            await asyncio.Event().wait()

    providers = ExcelProviders(dictionaries=Blocking())
    owned = []
    buffers = []

    def track_workbook():
        workbook = Workbook()
        close = workbook.close
        workbook.was_closed = False

        def close_tracked():
            workbook.was_closed = True
            close()

        workbook.close = close_tracked
        owned.append(workbook)
        return workbook

    def track_buffer():
        buffer = io.BytesIO()
        buffers.append(buffer)
        return buffer

    monkeypatch.setattr(writer_module, "Workbook", track_workbook)
    monkeypatch.setattr(writer_module, "io", SimpleNamespace(BytesIO=track_buffer))
    task = asyncio.create_task(
        ExcelWriter(settings()).write("数据", DictRow, [DictRow(value="zh")], providers=providers)
    )
    await started.wait()
    task.cancel()
    with pytest.raises(asyncio.CancelledError):
        await task
    assert owned[0].was_closed and buffers[0].closed

    started.clear()
    source = upload(xlsx([["字典"], ["中文"]]))
    source.file.seek(7)
    loaded = []
    original_load = load_workbook

    def track_load(*args, **kwargs):
        workbook = original_load(*args, **kwargs)
        loaded.append(workbook)
        return workbook

    monkeypatch.setattr(reader_module, "load_workbook", track_load)
    task = asyncio.create_task(ExcelReader(settings()).read(source, DictRow, providers=providers))
    await started.wait()
    task.cancel()
    with pytest.raises(asyncio.CancelledError):
        await task
    assert loaded[0]._archive.fp is None
    assert not source.file.closed and source.file.tell() == 7


@pytest.mark.asyncio
async def test_timezone_aware_datetime_is_not_silently_changed():
    class Row(BaseModel):
        at: Annotated[datetime, ExcelColumn("时间")]

    with pytest.raises(ExcelException) as raised:
        await ExcelWriter(settings()).write("数据", Row, [Row(at=datetime.now(timezone.utc))])
    assert raised.value.issues[0].field == "at"


@pytest.mark.asyncio
async def test_provider_custom_failure_keeps_position_and_cause():
    class DictRow(BaseModel):
        value: Annotated[str, ExcelColumn("字典", converter=DictConverter("language"))]

    failure = RuntimeError("real provider failure")

    class Broken:
        async def items(self, kind):
            raise failure

    providers = ExcelProviders(dictionaries=Broken())
    for operation in (
        ExcelReader(settings()).read(
            upload(xlsx([["字典"], ["中文"]])), DictRow, providers=providers
        ),
        ExcelWriter(settings()).write("数据", DictRow, [DictRow(value="zh")], providers=providers),
        ExcelWriter(settings()).template("模板", DictRow, providers=providers),
    ):
        with pytest.raises(ExcelException) as raised:
            await operation
        assert raised.value.__cause__ is failure
        assert raised.value.issues[0].column == 1 and raised.value.issues[0].field == "value"


def test_archive_actual_data_crc_duplicate_cells_and_nonseekable_stream():
    normal = xlsx([["文本"], ["hello"]])
    forged = bytearray(normal)
    position = forged.find(b"PK\x01\x02")
    # 中央目录谎报解压长度，ZipExtFile 的实际内容/CRC 校验必须拒绝。
    struct.pack_into("<L", forged, position + 24, 1)
    with pytest.raises(ExcelException) as raised:
        ExcelUploadValidator(settings()).validate(upload(forged))
    assert raised.value.__cause__ is not None
    duplicate = replace_member(
        normal, "xl/worksheets/sheet1.xml", lambda data: data.replace(b'r="A2"', b'r="A1"')
    )
    with pytest.raises(ExcelException) as raised:
        ExcelUploadValidator(settings()).validate(upload(duplicate))
    assert isinstance(raised.value.__cause__, ValueError)

    class NonSeekable(io.BytesIO):
        def seekable(self):
            return False

    source = upload(normal)
    source.file = NonSeekable(normal)
    with pytest.raises(ExcelException, match="定位"):
        ExcelUploadValidator(settings()).validate(source)
    assert not source.file.closed


@pytest.mark.asyncio
async def test_dropdowns_share_total_cell_budget_and_error_cleanup(monkeypatch):
    class OptionsRow(BaseModel):
        text: Annotated[str, ExcelColumn("文本", options=("a", "b", "c"))]

    with pytest.raises(ExcelException) as raised:
        await ExcelWriter(settings(max_cells=4)).write("数据", OptionsRow, [OptionsRow(text="a")])
    assert raised.value.error_code == ExcelErrorCodes.LIMIT
    closed = []
    original = Workbook.close

    def close(workbook):
        closed.append(workbook)
        return original(workbook)

    monkeypatch.setattr(Workbook, "close", close)
    with pytest.raises(ExcelException):
        await ExcelWriter(settings()).write("数据", TextRow, [TextRow(text="invalid\x01")])
    assert len(closed) == 1


@pytest.mark.asyncio
async def test_settings_and_unique_column_contracts():
    with pytest.raises(ValueError):
        settings(money_decimal_places=1)
    with pytest.raises(ValueError):
        settings(min_column_width=80, max_column_width=10)

    class Duplicate(BaseModel):
        first: Annotated[str, ExcelColumn("重复")]
        second: Annotated[str, ExcelColumn("重复")]

    with pytest.raises(ExcelException):
        ExcelSchema(Duplicate)


@pytest.mark.asyncio
async def test_operation_lookup_state_does_not_cross_applications():
    class DictRow(BaseModel):
        value: Annotated[str, ExcelColumn("字典", converter=DictConverter("language"))]

    first = ExcelProviders(dictionaries=Dictionaries("一"))
    second = ExcelProviders(dictionaries=Dictionaries("二"))
    outputs = await asyncio.gather(
        ExcelWriter(settings()).write("数据", DictRow, [DictRow(value="zh")], providers=first),
        ExcelWriter(settings()).write("数据", DictRow, [DictRow(value="zh")], providers=second),
    )
    for output, expected in zip(outputs, ("一", "二"), strict=True):
        workbook = load_workbook(output)
        try:
            assert workbook.active.cell(2, 1).value == expected
        finally:
            workbook.close()
            output.close()
