import json
import os
import re
from contextlib import asynccontextmanager
from datetime import datetime
from io import BytesIO
from pathlib import Path
from uuid import uuid4

import pytest
import pytest_asyncio
import yaml
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook
from pymysql.constants import CLIENT
from redis import Redis
from sqlalchemy import delete, func, insert, inspect, select
from sqlalchemy.engine import make_url
from sqlalchemy.ext.asyncio import create_async_engine

from fixtures.config_factory import ConfigFactory
from framework.starter_mq.model.outbox_job_parameters import OutboxJobParameters
from module_infra.dal.dataobject.job.job_do import JobDO
from module_system.dal.dataobject.dept.post_do import PostDO
from module_system.dal.dataobject.social.social_client_do import SocialClientDO
from server.starter_server import StarterServer

TARGETS_PATH = os.environ.get("DUSHAN_BUSINESS_DATABASE_TARGETS")
TARGETS = (
    json.loads(Path(TARGETS_PATH).read_text(encoding="utf-8-sig"))
    if TARGETS_PATH
    else [{"name": "unconfigured", "url": ""}]
)
SHARED_TARGETS = [row for row in TARGETS if row["name"] != "dm"]
DM_TARGETS = [row for row in TARGETS if row["name"] == "dm"]


@pytest.mark.asyncio(loop_scope="class")
class TestMultidbBusiness:
    __test__ = bool(SHARED_TARGETS)
    MYSQL_DATABASES = {"mysql", "tidb", "oceanbase"}

    @pytest_asyncio.fixture(
        scope="class",
        loop_scope="class",
        params=SHARED_TARGETS,
        ids=[row["name"] for row in SHARED_TARGETS],
    )
    @classmethod
    async def business_database(cls, request, tmp_path_factory):
        async with cls._business_database(request, tmp_path_factory) as database:
            yield database

    @classmethod
    @asynccontextmanager
    async def _business_database(cls, request, tmp_path_factory):
        """为每个显式提供的实例创建独立空库，装载发布 SQL 后在退出时删除。"""
        if not TARGETS_PATH:
            pytest.skip("需要数据库实例清单 DUSHAN_BUSINESS_DATABASE_TARGETS")
        root = Path(__file__).resolve().parents[2]
        name = "multidb_" + uuid4().hex
        target = request.param
        url = make_url(target["url"])
        mysql = target["name"] in cls.MYSQL_DATABASES
        dm = target["name"] == "dm"
        if dm:
            name = name.upper()
        options = {"client_flag": CLIENT.MULTI_STATEMENTS} if mysql else {}
        admin_url = (
            url
            if dm
            else url.set(
                database="mysql"
                if mysql
                else "test"
                if target["name"] == "kingbase"
                else "postgres"
            )
        )
        admin = create_async_engine(admin_url, isolation_level="AUTOCOMMIT")
        quoted = f"`{name}`" if mysql else f'"{name}"'
        created = False
        engine = None
        loaded = []
        evidence_path = tmp_path_factory.mktemp("business-load") / f"{target['name']}.json"
        evidence = {
            "database": target["name"],
            "schema": name,
            "test": request.node.nodeid,
            "version": None,
            "loaded": loaded,
            "dropped": False,
        }
        evidence_path.write_text(
            json.dumps(evidence, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        try:
            async with admin.connect() as connection:
                if dm:
                    password = "Business@" + uuid4().hex
                    await connection.exec_driver_sql(
                        f'CREATE USER {quoted} IDENTIFIED BY "{password}"'
                    )
                    created = True
                    await connection.exec_driver_sql(f"GRANT RESOURCE TO {quoted}")
                    test_url = url.set(username=name, password=password, database="")
                else:
                    suffix = " CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci" if mysql else ""
                    if target["name"] == "opengauss":
                        suffix = " DBCOMPATIBILITY 'PG'"
                    await connection.exec_driver_sql(f"CREATE DATABASE {quoted}{suffix}")
                    created = True
                    test_url = url.set(database=name)
            engine = create_async_engine(
                test_url, isolation_level="AUTOCOMMIT", connect_args=options
            )
            sql_root = root / "dushan-admin-backend/sql" / target["name"]
            paths = [
                path
                for module in ("system", "infra")
                for path in sorted((sql_root / module).glob("*.sql"))
            ]
            assert len(paths) == 8, sql_root
            async with engine.connect() as connection:
                if dm:
                    evidence["version"] = (
                        await connection.exec_driver_sql("SELECT BANNER FROM V$VERSION")
                    ).scalar()
                else:
                    evidence["version"] = await connection.scalar(select(func.version()))
                raw = await connection.get_raw_connection()
                driver = raw.driver_connection
                for path in paths:
                    sql = path.read_text(encoding="utf-8")
                    if mysql:
                        async with driver.cursor() as cursor:
                            await cursor.execute(sql)
                            while await cursor.nextset():
                                pass
                    elif dm:
                        for statement in cls._dm_statements(sql):
                            await connection.exec_driver_sql(statement)
                    else:
                        await driver.execute(sql)
                    loaded.append(path.relative_to(root).as_posix())
            yield target["name"], test_url.render_as_string(hide_password=False), engine, loaded
        finally:
            try:
                if engine is not None:
                    await engine.dispose()
                if created:
                    async with admin.connect() as connection:
                        await connection.exec_driver_sql(
                            f"DROP USER {quoted} CASCADE" if dm else f"DROP DATABASE {quoted}"
                        )
                    evidence["dropped"] = True
            finally:
                await admin.dispose()
                evidence_path.write_text(
                    json.dumps(evidence, ensure_ascii=False, indent=2), encoding="utf-8"
                )

    @pytest_asyncio.fixture(scope="class", loop_scope="class")
    @classmethod
    async def business_app(cls, business_database, tmp_path_factory):
        async with cls._business_app(business_database, tmp_path_factory) as app:
            yield app

    @classmethod
    @asynccontextmanager
    async def _business_app(cls, business_database, tmp_path_factory):
        """启动 System 与 Infra 真实应用，Redis 使用任务独占的第 13 个逻辑库。"""
        resource_path = os.environ.get("DUSHAN_SYSTEM_RESOURCES")
        if resource_path is None:
            pytest.skip("需要 Redis 资源清单 DUSHAN_SYSTEM_RESOURCES")
        resources = json.loads(Path(resource_path).read_text(encoding="utf-8-sig"))
        _, url, _, _ = business_database
        folder = tmp_path_factory.mktemp("multidb-business")
        values = ConfigFactory.values()
        values["modules"]["enabled"] = ["framework", "system", "infra"]
        values["server"]["reload"] = False
        values["banner"]["enabled"] = False
        values["log"].update(enable_file_overall=False, console_level="WARNING")
        values["page"]["fetch_all_enabled"] = True
        models = values["config"]["models"]
        models["database"].update(
            enabled=True,
            health_check_enabled=False,
            dynamic_enabled=True,
            id_strategy="snowflake",
            snowflake_machine_id=936,
            sources=[
                dict(name="primary", url=url, role="primary", weight=100, pool=None, tls=None)
            ],
        )
        models["cache"].update(
            enabled=True,
            host="127.0.0.1",
            port=resources["redis_port"],
            clients=[dict(name="default", db=13)],
        )
        models["security"].update(enabled=True, permission_cache_enabled=True)
        models["data_permission"].update(enabled=True, cache_enabled=True)
        models["protection"]["enabled"] = True
        models["qr_login"]["enabled"] = False
        models["system"].update(
            workload_credential=uuid4().hex + uuid4().hex,
            message_signing_key=uuid4().hex + uuid4().hex,
            sms_callback_token=uuid4().hex + uuid4().hex,
            refresh_cookie_secure=False,
            allowed_origins=["http://testserver"],
        )
        (folder / "application.yaml").write_text(
            yaml.safe_dump(values, allow_unicode=True), encoding="utf-8"
        )
        with Redis(host="127.0.0.1", port=resources["redis_port"], db=13) as redis:
            redis.flushdb()
            try:
                app = StarterServer.create_app(base_dir=folder, app_env="dev", environ={})
                async with app.router.lifespan_context(app):
                    yield app
            finally:
                redis.flushdb()

    @pytest_asyncio.fixture(scope="class", loop_scope="class")
    @classmethod
    async def business_client(cls, business_app):
        async with cls._business_client(business_app) as client:
            yield client

    @classmethod
    @asynccontextmanager
    async def _business_client(cls, business_app):
        """用种子管理员登录并保留服务端签发的令牌。"""
        async with AsyncClient(
            transport=ASGITransport(app=business_app), base_url="http://testserver"
        ) as client:
            result = cls._data(
                await client.post(
                    "/admin-api/system/auth/login",
                    json={"username": "admin", "password": "admin123"},
                    headers={"X-Tenant-Id": "1"},
                )
            )
            client.headers["Authorization"] = "Bearer " + result["accessToken"]
            yield client

    @staticmethod
    def _dm_statements(source):
        """分割发布 SQL 的语句，保留字符串和注释内的分号。"""
        start = 0
        for token in re.finditer(
            r"'(?:''|[^'])*'|\"(?:\"\"|[^\"])*\"|--[^\n]*|/\*.*?\*/|;", source, re.DOTALL
        ):
            if token.group() == ";":
                yield source[start : token.start()]
                start = token.end()
        assert not source[start:].strip(), "发布 SQL 最后一条语句必须以分号结尾"

    @staticmethod
    def _data(response):
        """校验真实 HTTP 与业务响应后返回数据。"""
        assert response.status_code == 200, response.text
        payload = response.json()
        assert payload["code"] == 0, (response.request.url.path, payload)
        return payload["data"]

    @classmethod
    async def _create(cls, client, resource, payload):
        """通过业务创建接口写入并取得主键。"""
        return cls._data(await client.post(f"/admin-api/system/{resource}/create", json=payload))

    async def test_load_all_sql(self, business_database):
        """确认八个发布文件均已成功执行，并核对 System 与 Infra 表实际存在。"""
        name, _, engine, loaded = business_database
        assert len(loaded) == 8
        assert all(path.startswith(f"dushan-admin-backend/sql/{name}/") for path in loaded)
        async with engine.connect() as connection:
            tables = await connection.run_sync(lambda sync: inspect(sync).get_table_names())
            if name == "dm":
                for table in ("system_tenant", "system_users", "infra_config_type"):
                    result = await connection.exec_driver_sql(
                        f"SELECT IDENT_CURRENT('{table.upper()}') + IDENT_INCR('{table.upper()}'), "
                        f"(SELECT max(id) FROM {table})"
                    )
                    next_id, maximum_id = result.one()
                    assert next_id > maximum_id, table
            elif name not in self.MYSQL_DATABASES:
                for table in ("system_tenant", "system_users", "infra_config_type"):
                    result = await connection.exec_driver_sql(
                        f"SELECT nextval(pg_get_serial_sequence('{table}', 'id')) > "
                        f"(SELECT max(id) FROM {table})"
                    )
                    assert result.scalar_one() is True, table
        assert {"system_users", "system_post", "infra_config_data", "infra_job"} <= set(tables)
        assert len(tables) == 54

    async def test_outbox_dispatch_job_seed(self, business_database):
        """补扫任务按租户扇出，使用处理器默认参数，等待开启 MQ 后启用。"""
        _, _, engine, _ = business_database
        async with engine.connect() as connection:
            result = await connection.execute(
                select(JobDO.__table__).where(JobDO.handler_name == "mq.outbox.dispatch")
            )
            row = result.mappings().one()
        assert row["status"] == 2
        assert "开启 MQ 后启用" in row["name"]
        assert row["deleted"] is False
        assert row["fan_out"] is True
        assert row["tenant_id"] is None
        assert row["cron_expression"] == "* * * * *"
        assert row["parameters"] == OutboxJobParameters().model_dump()
        assert json.loads(row["handler_param"]) == row["parameters"]
        assert row["max_instances"] == 1
        assert row["monitor_timeout"] is None
        assert row["timeout_seconds"] == 300
        assert row["retry_count"] == 3
        assert row["retry_interval"] == 5000
        assert row["retry_backoff"] == 1
        assert row["stop_after_failure"] is False
        assert row["revision"]
        assert row["effective_at"] is not None

    async def test_json_default(self, business_database):
        """省略认证配置执行真实模型插入，确认默认值保存为空 JSON 对象。"""
        _, _, engine, _ = business_database
        identifier = uuid4().int >> 65
        now = datetime(2026, 10, 5)
        async with engine.begin() as connection:
            await connection.execute(
                insert(SocialClientDO).values(
                    id=identifier,
                    name="JSON默认值验证",
                    social_type=31,
                    user_type=2,
                    client_id=uuid4().hex,
                    client_secret="test-only",
                    tenant_id="1",
                    creator="system",
                    updater="system",
                    create_time=now,
                    update_time=now,
                    deleted=False,
                    status=1,
                )
            )
            actual = await connection.scalar(
                select(SocialClientDO.auth_config).where(SocialClientDO.id == identifier)
            )
            assert actual == {}
            await connection.execute(delete(SocialClientDO).where(SocialClientDO.id == identifier))

    async def test_admin_login(self, business_client):
        """管理员令牌可以取得本人及权限信息。"""
        data = self._data(await business_client.get("/admin-api/system/auth/get-permission-info"))
        assert data["user"]["username"] == "admin"

    @pytest.mark.parametrize(
        "website, expected", [("http://localhost", "1"), ("https://missing.example", None)]
    )
    async def test_tenant_website_json_lookup(self, business_client, website, expected):
        """租户域名的 JSON 数组查询保持精确命中与未命中的相同语义。"""
        data = self._data(
            await business_client.get(
                "/admin-api/system/tenant/get-by-website", params={"website": website}
            )
        )
        assert (data["id"] if data is not None else None) == expected

    @pytest.mark.parametrize(
        "channels, expected_indices",
        [
            (["INTERNAL"], [0, 2]),
            (["SMS"], []),
            (["INTERNAL", "MAIL"], [0, 1, 2]),
        ],
        ids=["matching", "missing", "or"],
    )
    async def test_notice_channel_json_filter(self, business_client, channels, expected_indices):
        """通知模板按渠道精确筛选，多渠道取并集且同一模板不重复计数。"""
        title = "channels" + uuid4().hex[:12]
        identifiers = []
        for saved_channels in (["INTERNAL"], ["MAIL"], ["INTERNAL", "MAIL"]):
            identifiers.append(
                await self._create(
                    business_client,
                    "notification",
                    {
                        "title": title,
                        "content": "渠道筛选验证",
                        "type": 2,
                        "userType": 2,
                        "channels": saved_channels,
                        "status": 1,
                        "publisher": "渠道筛选测试",
                    },
                )
            )
        data = self._data(
            await business_client.get(
                "/admin-api/system/notification/page",
                params={"title": title, "channels": channels},
            )
        )
        assert {row["id"] for row in data["items"]} == {
            identifiers[index] for index in expected_indices
        }
        assert data["total"] == len(expected_indices)
        for identifier in identifiers:
            self._data(
                await business_client.delete(
                    "/admin-api/system/notification/delete", params={"id": identifier}
                )
            )

    async def test_demo_login_and_write_rejected(self, business_app):
        """演示用户可以登录及查询，写请求被只读保护拒绝。"""
        async with AsyncClient(
            transport=ASGITransport(app=business_app), base_url="http://testserver"
        ) as client:
            data = self._data(
                await client.post(
                    "/admin-api/system/auth/login",
                    json={"username": "demo", "password": "demo123456"},
                    headers={"X-Tenant-Id": "1"},
                )
            )
            client.headers["Authorization"] = "Bearer " + data["accessToken"]
            info = self._data(await client.get("/admin-api/system/auth/get-permission-info"))
            assert info["user"]["username"] == "demo"
            self._data(await client.get("/admin-api/system/user/page"))
            denied = await client.post(
                "/admin-api/system/dept/create",
                json={"name": "denied", "parentId": "0", "sort": 1, "status": 1},
            )
            assert denied.json()["code"] == 901, denied.text

    @pytest.mark.parametrize(
        "path",
        [
            "system/user/page",
            "system/dept/list",
            "system/dept/post/page",
            "system/permission/role/page",
            "system/permission/menu/list",
            "system/dict/type/page",
            "system/dict/data/page",
            "system/tenant/page",
            "system/tenant/package/page",
            "system/logger/login-log/page",
            "system/logger/operate-log/page",
            "system/mail/account/page",
            "system/mail/template/page",
            "system/mail/log/page",
            "system/sms/channel/page",
            "system/sms/template/page",
            "system/sms/log/page",
            "system/social/client/page",
            "system/social/user/page",
            "system/oauth2/client/page",
            "system/oauth2/token/page",
            "system/notification/page",
            "system/announcement/page",
            "infra/config/page",
            "infra/config/type/page",
            "infra/file/config/page",
            "infra/file/page",
            "infra/data-source/page",
            "infra/job/page",
            "infra/job/log/page",
            "infra/mq/page",
            "infra/mq/log/page",
            "infra/logger/api-access-log/page",
            "infra/logger/api-error-log/page",
            "infra/codegen/table/page",
        ],
    )
    async def test_list_and_page(self, business_client, path):
        """逐项执行 System 与 Infra 主要列表和分页接口。"""
        data = self._data(await business_client.get("/admin-api/" + path))
        if path.endswith("/page"):
            assert isinstance(data["items"], list)
            assert data["total"] >= len(data["items"])
        else:
            assert isinstance(data, list)

    @pytest.mark.parametrize("resource", ["dept", "dept/post", "dict/type", "dict/data"])
    async def test_create_update_delete(self, business_client, resource):
        """部门、岗位、字典类型及字典数据完成创建、修改和逻辑删除。"""
        key = "crud_" + uuid4().hex[:12]
        payload = {"name": key, "status": 1}
        parent = None
        match resource:
            case "dept":
                payload.update(parentId="0", sort=1)
            case "dept/post":
                payload.update(code=key, sort=1)
            case "dict/type":
                payload.update(type=key)
            case "dict/data":
                parent = await self._create(
                    business_client, "dict/type", {"name": key, "type": key, "status": 1}
                )
                payload = {"label": key, "value": key, "dictType": key, "sort": 1, "status": 1}
        path = f"/admin-api/system/{resource}"
        identifier = await self._create(business_client, resource, payload)
        field = "label" if resource == "dict/data" else "name"
        payload.update(id=identifier)
        payload[field] = key + "改"
        self._data(await business_client.put(path + "/update", json=payload))
        data = self._data(await business_client.get(path + "/get", params={"id": identifier}))
        assert data[field] == key + "改"
        self._data(await business_client.delete(path + "/delete", params={"id": identifier}))
        endpoint = "/list" if resource == "dept" else "/page"
        data = self._data(
            await business_client.get(path + endpoint, params={field: payload[field]})
        )
        rows = data if resource == "dept" else data["items"]
        assert identifier not in {row["id"] for row in rows}
        if parent is not None:
            self._data(
                await business_client.delete(
                    "/admin-api/system/dict/type/delete", params={"id": parent}
                )
            )

    @pytest.mark.parametrize("symbol", ["%", "_", "\\"])
    @pytest.mark.parametrize("resource", ["dept", "dept/post", "dict/type"])
    async def test_like_special_characters(self, business_client, resource, symbol):
        """模糊查询中的百分号、下划线及反斜线按字面量匹配。"""
        key = "like_" + uuid4().hex[:8]
        identifiers = []
        for suffix in (symbol, "X"):
            payload = {"name": key + suffix + "尾", "status": 1}
            if resource == "dept":
                payload.update(parentId="0", sort=1)
            elif resource == "dept/post":
                payload.update(code=uuid4().hex, sort=1)
            else:
                payload.update(type=uuid4().hex)
            identifiers.append(await self._create(business_client, resource, payload))
        path = f"/admin-api/system/{resource}"
        endpoint = "/list" if resource == "dept" else "/page"
        data = self._data(await business_client.get(path + endpoint, params={"name": key + symbol}))
        rows = data if resource == "dept" else data["items"]
        assert [row["id"] for row in rows] == identifiers[:1]
        for identifier in identifiers:
            self._data(await business_client.delete(path + "/delete", params={"id": identifier}))

    @pytest.mark.parametrize("path", ["system/dept/post", "system/dict/type", "infra/config"])
    async def test_export_excel(self, business_client, path):
        """读取实际导出的 XLSX 工作簿，验证查询、序列化与响应输出。"""
        response = await business_client.get(f"/admin-api/{path}/export-excel")
        assert response.status_code == 200, response.text
        assert "spreadsheetml" in response.headers["content-type"]
        workbook = load_workbook(BytesIO(response.content), read_only=True)
        try:
            assert workbook.active.max_row >= 1
            assert next(workbook.active.values)
        finally:
            workbook.close()

    async def test_soft_delete_reuses_unique_code(self, business_client, business_database):
        """连续删除及复用岗位编码，核对唯一约束忽略全部已删除记录。"""
        _, _, engine, _ = business_database
        code = "reuse_" + uuid4().hex
        payload = {"name": code, "code": code, "sort": 1, "status": 1}
        identifiers = []
        for _ in range(3):
            identifier = await self._create(business_client, "dept/post", payload)
            identifiers.append(int(identifier))
            self._data(
                await business_client.delete(
                    "/admin-api/system/dept/post/delete", params={"id": identifier}
                )
            )
        async with engine.connect() as connection:
            deleted = await connection.scalar(
                select(func.count())
                .select_from(PostDO)
                .where(PostDO.id.in_(identifiers), PostDO.deleted.is_(True))
            )
        assert deleted == 3


class TestDmBusiness(TestMultidbBusiness):
    __test__ = bool(DM_TARGETS)

    @pytest_asyncio.fixture(
        loop_scope="class", params=DM_TARGETS, ids=[row["name"] for row in DM_TARGETS]
    )
    @classmethod
    async def business_database(cls, request, tmp_path_factory):
        """达梦每个测试使用独立用户及同名模式，退出时连同数据一起删除。"""
        async with cls._business_database(request, tmp_path_factory) as database:
            yield database

    @pytest_asyncio.fixture(loop_scope="class")
    @classmethod
    async def business_app(cls, business_database, tmp_path_factory):
        async with cls._business_app(business_database, tmp_path_factory) as app:
            yield app

    @pytest_asyncio.fixture(loop_scope="class")
    @classmethod
    async def business_client(cls, business_app):
        async with cls._business_client(business_app) as client:
            yield client
