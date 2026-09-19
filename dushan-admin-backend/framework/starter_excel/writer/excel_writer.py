import asyncio
import csv
import io
import math
from collections.abc import Sequence
from datetime import date, datetime, time
from decimal import Decimal
from typing import TypeVar

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.utils.exceptions import IllegalCharacterError
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import BaseModel

from framework.common.page.schemas.page_query import PageQuery
from framework.common.utils.asyncio_utils import AsyncioUtils
from framework.starter_di.decorators.components import framework
from framework.starter_excel.config.excel_settings import ExcelSettings
from framework.starter_excel.converter.dict_converter import DictConverter
from framework.starter_excel.converter.enum_converter import EnumConverter
from framework.starter_excel.core.excel_schema import ExcelSchema
from framework.starter_excel.definitions.constants.excel_error_codes import ExcelErrorCodes
from framework.starter_excel.exception.excel_exception import ExcelException
from framework.starter_excel.model.conversion_context import ConversionContext
from framework.starter_excel.model.excel_issue import ExcelIssue
from framework.starter_excel.model.excel_providers import ExcelProviders

T = TypeVar("T", bound=BaseModel)


@framework
class ExcelWriter:
    """生成 XLSX 内存流；成功后流归调用方，失败和取消会关闭本次所有资源。"""

    def __init__(self, settings: ExcelSettings) -> None:
        self.settings = settings

    def enable_export_fetch_all(self, page: PageQuery) -> None:
        page.enable_fetch_all(max_rows=self.settings.max_export_rows)

    @staticmethod
    def export_fields(model: type[BaseModel]) -> list[dict[str, str]]:
        return ExcelSchema(model).export_fields()

    async def template(
        self, sheet_name: str, model: type[T], *, providers: ExcelProviders | None = None
    ) -> io.BytesIO:
        """模板包含相同表头、样式与真实下拉选项，不虚构一行示例数据。"""
        return await self.write(sheet_name, model, [], providers=providers)

    async def write(
        self,
        sheet_name: str,
        model: type[T],
        data: Sequence[T],
        *,
        fields: list[str] | None = None,
        providers: ExcelProviders | None = None,
    ) -> io.BytesIO:
        columns = ExcelSchema(model).export_columns(fields)
        if len(data) > self.settings.max_export_rows or len(columns) > self.settings.max_columns:
            raise ExcelException(ExcelErrorCodes.LIMIT, "导出行列数超过限制")
        if (len(data) + 1) * len(columns) > self.settings.max_cells:
            raise ExcelException(ExcelErrorCodes.LIMIT, "导出单元格数超过限制")
        context = ConversionContext(self.settings, providers or ExcelProviders())
        output = io.BytesIO()
        workbook = Workbook()
        success = False
        try:
            sheet = workbook.active
            sheet.title = sheet_name
            sheet.freeze_panes = "A2"
            total_text = 0
            widths = [0] * len(columns)
            for index, column in enumerate(columns.values(), start=1):
                total_text += self._cell(sheet, 1, index, column.title, total_text)
                cell = sheet.cell(1, index)
                cell.font = Font(
                    bold=self.settings.header_font_bold, size=self.settings.header_font_size
                )
                cell.alignment = Alignment(wrap_text=True)
                widths[index - 1] = len(column.title)
            for row_number, item in enumerate(data, start=2):
                if row_number % 100 == 0:
                    await asyncio.sleep(0)
                # 只序列化无转换器字段；Area 含父子关系，不能递归 model_dump。
                scalar_fields = {
                    name for name, column in columns.items() if column.converter is None
                }
                serialized = item.model_dump(mode="python", include=scalar_fields, by_alias=False)
                for index, (name, column) in enumerate(columns.items(), start=1):
                    try:
                        if column.converter is None:
                            value = serialized[name]
                        else:
                            value = getattr(item, name)
                            if value is not None:
                                value = await column.converter.to_excel(value, context)
                        total_text += self._cell(sheet, row_number, index, value, total_text)
                        if column.number_format is not None:
                            sheet.cell(row_number, index).number_format = column.number_format
                        widths[index - 1] = max(
                            widths[index - 1], len(str(sheet.cell(row_number, index).value or ""))
                        )
                    except (
                        ValueError,
                        TypeError,
                        ArithmeticError,
                        csv.Error,
                        IllegalCharacterError,
                    ) as exc:
                        raise ExcelException(
                            ExcelErrorCodes.CONVERSION,
                            "Excel 导出字段转换失败",
                            issues=[
                                ExcelIssue(row_number, index, name, "字段转换或单元格内容无效")
                            ],
                            cause=exc,
                        ) from exc
                    except ExcelException:
                        raise
                    except Exception as exc:
                        # Provider 故障类型由业务定义，保留 cause 并定位当前单元格。
                        raise ExcelException(
                            ExcelErrorCodes.CONVERSION,
                            "Excel 外部数据查询失败",
                            issues=[ExcelIssue(row_number, index, name, "外部数据查询失败")],
                            cause=exc,
                        ) from exc
            await self._dropdowns(sheet, columns, len(data), context, total_text)
            if self.settings.auto_adjust_column_width:
                for index, width in enumerate(widths, start=1):
                    sheet.column_dimensions[get_column_letter(index)].width = min(
                        self.settings.max_column_width,
                        max(self.settings.min_column_width, width + 3),
                    )
            await asyncio.sleep(0)
            await AsyncioUtils.run_cancellation_shielded(asyncio.to_thread(workbook.save, output))
            output.seek(0)
            success = True
            return output
        except (OSError, ValueError, KeyError, TypeError, IllegalCharacterError) as exc:
            raise ExcelException(ExcelErrorCodes.WRITE, "写入 XLSX 工作簿失败", cause=exc) from exc
        finally:
            workbook.close()
            if not success:
                output.close()

    def _cell(self, sheet, row: int, column: int, value, text_bytes: int) -> int:
        if isinstance(value, (datetime, time)) and value.tzinfo is not None:
            raise ValueError("Excel 日期不包含时区，请显式转换为无时区日期")
        if isinstance(value, float) and not math.isfinite(value):
            raise ValueError("Excel 不支持非有限数字")
        if isinstance(value, Decimal):
            if not value.is_finite():
                raise ValueError("Excel 不支持非有限数字")
            value = str(value)
        if type(value) is int and abs(value) >= 10**15:
            value = str(value)
        if value is not None and not isinstance(
            value, (str, int, float, bool, date, datetime, time)
        ):
            raise ValueError("复杂字段必须声明转换器")
        size = 0
        if isinstance(value, str):
            size = len(value.encode("utf-8"))
            if (
                len(value) > self.settings.max_cell_text_length
                or text_bytes + size > self.settings.max_uncompressed_size_bytes
            ):
                raise ExcelException(ExcelErrorCodes.LIMIT, "导出文本长度或总字节数超过限制")
        cell = sheet.cell(row, column, value)
        if isinstance(value, str):
            # 所有外部文本都写为字符串，包括表头、下拉选项和带空白的公式前缀。
            cell.data_type = "s"
        return size

    async def _dropdowns(self, sheet, columns, data_count, context, text_bytes) -> None:
        options_sheet = None
        option_cells = (data_count + 1) * len(columns)
        for index, (name, column) in enumerate(columns.items(), start=1):
            options = column.options
            if options is None and isinstance(column.converter, (DictConverter, EnumConverter)):
                try:
                    options = await column.converter.options(context)
                except Exception as exc:
                    raise ExcelException(
                        ExcelErrorCodes.CONVERSION,
                        "Excel 下拉选项查询失败",
                        issues=[ExcelIssue(1, index, name, "下拉选项查询失败")],
                        cause=exc,
                    ) from exc
            if not options:
                continue
            option_cells += len(options)
            if (
                len(options) > self.settings.max_dropdown_options
                or option_cells > self.settings.max_cells
            ):
                raise ExcelException(ExcelErrorCodes.LIMIT, "下拉选项数超过限制")
            if options_sheet is None:
                options_sheet = sheet.parent.create_sheet("_excel_options")
                options_sheet.sheet_state = "hidden"
            for row, option in enumerate(options, start=1):
                text_bytes += self._cell(options_sheet, row, index, option, text_bytes)
            letter = get_column_letter(index)
            range_name = f"_excel_options_{index}"
            reference = (
                f"{quote_sheetname(options_sheet.title)}!${letter}$1:${letter}${len(options)}"
            )
            sheet.parent.defined_names.add(DefinedName(range_name, attr_text=reference))
            validation = DataValidation(
                type="list",
                formula1=range_name,
                allow_blank=True,
                showDropDown=False,
                showErrorMessage=True,
                errorTitle="输入错误",
                error="请使用下拉选项中的值",
            )
            validation.add(f"{letter}2:{letter}{max(data_count, self.settings.template_rows) + 1}")
            sheet.add_data_validation(validation)
