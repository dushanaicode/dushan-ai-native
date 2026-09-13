import asyncio
import csv
from functools import partial
from itertools import islice
from typing import TypeVar
from zipfile import BadZipFile

from openpyxl import load_workbook
from pydantic import BaseModel, ValidationError

from framework.common.utils.asyncio.asyncio_utils import AsyncioUtils
from framework.common.utils.asyncio.cleanup_utils import CleanupUtils
from framework.starter_di.decorators.components import framework
from framework.starter_di.enums.component_scope_enum import ComponentScopeEnum
from framework.starter_excel.config.excel_settings import ExcelSettings
from framework.starter_excel.core.excel_schema import ExcelSchema
from framework.starter_excel.exception.excel_error import ExcelError
from framework.starter_excel.exception.excel_error_codes import ExcelErrorCodes
from framework.starter_excel.handler.excel_upload_validator import ExcelUploadValidator
from framework.starter_excel.model.conversion_context import ConversionContext
from framework.starter_excel.model.excel_issue import ExcelIssue
from framework.starter_excel.model.excel_providers import ExcelProviders
from framework.starter_excel.model.excel_upload import ExcelUpload

T = TypeVar("T", bound=BaseModel)


@framework(scope=ComponentScopeEnum.SINGLETON)
class ExcelReader:
    """按表头导入 XLSX；整批成功才返回对象，最多保留配置数量的行列错误。"""

    def __init__(self, settings: ExcelSettings) -> None:
        self.settings = settings
        self._import_slots = asyncio.Semaphore(settings.max_concurrent_imports)

    async def read(
        self,
        upload: ExcelUpload,
        model: type[T],
        *,
        providers: ExcelProviders | None = None,
        sheet_name: str | None = None,
    ) -> list[T]:
        """默认读取活动表；取消等待文件线程终态和清理后传播，上传流仍归调用方。"""
        schema = ExcelSchema(model)
        context = ConversionContext(self.settings, providers or ExcelProviders())
        # Native DI 每个应用只持有一个 Reader；额度持续到文件线程和清理均已结束。
        async with self._import_slots:
            if not upload.file.seekable():
                raise ExcelError(ExcelErrorCodes.VALIDATION, "XLSX 上传流必须支持定位")
            original = upload.file.tell()
            worker = asyncio.create_task(asyncio.to_thread(self._load_workbook, upload))
            workbook = rows = primary = None
            try:
                try:
                    workbook = await AsyncioUtils.run_cancellation_shielded(worker)
                except asyncio.CancelledError:
                    # 加载可能已经成功；取消不能丢掉刚创建的工作簿而跳过关闭。
                    if not worker.cancelled() and worker.exception() is None:
                        workbook = worker.result()
                    raise
                if sheet_name is not None and sheet_name not in workbook.sheetnames:
                    raise ExcelError(ExcelErrorCodes.VALIDATION, "指定的工作表不存在")
                sheet = workbook.active if sheet_name is None else workbook[sheet_name]
                if sheet is None or sheet not in workbook.worksheets:
                    raise ExcelError(ExcelErrorCodes.VALIDATION, "没有可读取的数据工作表")
                # 不能信任上传者填写的 dimension；真实坐标已由 XML 预检约束。
                sheet.reset_dimensions()
                rows = sheet.iter_rows()
                return await self._rows(rows, schema, context)
            except (OSError, BadZipFile, ValueError, KeyError) as exc:
                primary = ExcelError(ExcelErrorCodes.READ, "读取 XLSX 工作簿失败", cause=exc)
            except BaseException as exc:
                primary = exc
            finally:
                error, cancellation = await CleanupUtils.run_cancellation_safe_cleanup(
                    partial(asyncio.to_thread, self._close, workbook, rows, upload, original),
                    "关闭 Excel 导入资源",
                )
                CleanupUtils.raise_collected_cleanup_errors(
                    "Excel 导入清理失败",
                    [] if error is None else [error],
                    primary_error=primary,
                    caller_cancellation=cancellation,
                )

    def _load_workbook(self, upload: ExcelUpload):
        """只在线程中执行文件工作，业务 Provider 和模型回调留在应用事件循环。"""
        ExcelUploadValidator(self.settings).validate(upload)
        upload.file.seek(0)
        return load_workbook(upload.file, read_only=True, data_only=False, keep_links=False)

    @staticmethod
    def _close(workbook, rows, upload: ExcelUpload, original: int) -> None:
        actions = []
        if rows is not None:
            actions.append(rows.close)
        if workbook is not None:
            actions.append(workbook.close)
        actions.append(partial(upload.file.seek, original))
        errors = []
        for action in actions:
            try:
                action()
            except BaseException as exc:
                errors.append(exc)
        CleanupUtils.raise_collected_cleanup_errors("Excel 文件关闭失败", errors)

    async def _rows(self, rows, schema: ExcelSchema, context: ConversionContext) -> list:
        header = await AsyncioUtils.run_cancellation_shielded(asyncio.to_thread(next, rows, ()))
        mapping = self._mapping(header, schema)
        results: list[BaseModel] = []
        issues: list[ExcelIssue] = []
        first_cause = None
        cells = len(header)
        row_number = 1
        # 每次只物化 100 行，避免在线程中一次保留整个表或执行用户回调。
        while batch := await AsyncioUtils.run_cancellation_shielded(
            asyncio.to_thread(list, islice(rows, 100))
        ):
            for row in batch:
                row_number += 1
                cells += len(row)
                if (
                    row_number > self.settings.max_import_rows + 1
                    or cells > self.settings.max_cells
                ):
                    raise ExcelError(ExcelErrorCodes.LIMIT, "导入实际行数或迭代单元格数超过限制")
                values = {}
                has_value = False
                row_issues = len(issues)
                for index, name in mapping.items():
                    cell = row[index - 1] if index <= len(row) else None
                    value = cell.value if cell is not None else None
                    column = schema.columns[name]
                    if isinstance(value, str):
                        value = value.strip() if column.trim else value
                    if value == "":
                        value = None
                    has_value |= value is not None
                    try:
                        if (
                            cell is not None
                            and cell.data_type == "f"
                            and self.settings.formula_policy == "reject"
                        ):
                            raise ValueError("不允许导入公式单元格")
                        if cell is not None and cell.data_type == "e":
                            raise ValueError("Excel 单元格包含错误值")
                        if (
                            isinstance(value, str)
                            and len(value) > self.settings.max_cell_text_length
                        ):
                            raise ValueError("单元格文本长度超过限制")
                        if value is not None and column.converter is not None:
                            value = await column.converter.to_python(value, context)
                        values[name] = schema.normalize(name, value)
                    except (ValueError, TypeError, ArithmeticError, csv.Error) as exc:
                        first_cause = first_cause or exc
                        issues.append(
                            ExcelIssue(row_number, index, name, "单元格类型、公式或转换结果无效")
                        )
                    except ExcelError:
                        raise
                    except Exception as exc:
                        # SPI 可抛出自己的异常类型；仅补行列位置，原始故障和取消不丢失。
                        raise ExcelError(
                            ExcelErrorCodes.CONVERSION,
                            "Excel 外部数据查询失败",
                            issues=[ExcelIssue(row_number, index, name, "外部数据查询失败")],
                            cause=exc,
                        ) from exc
                    if len(issues) >= self.settings.max_errors:
                        self._raise_issues(issues, first_cause)
                if not has_value or len(issues) != row_issues:
                    continue
                try:
                    results.append(
                        schema.model.model_validate(values, by_alias=False, by_name=True)
                    )
                except ValidationError as exc:
                    first_cause = first_cause or exc
                    reverse = {name: index for index, name in mapping.items()}
                    for error in exc.errors(
                        include_input=False, include_context=False, include_url=False
                    ):
                        name = str(error["loc"][0]) if error["loc"] else ""
                        issues.append(
                            ExcelIssue(
                                row_number,
                                reverse.get(name),
                                name,
                                f"字段约束不符：{error['type']}",
                            )
                        )
                        if len(issues) >= self.settings.max_errors:
                            self._raise_issues(issues, first_cause)
        if issues:
            self._raise_issues(issues, first_cause)
        return results

    def _mapping(self, header, schema: ExcelSchema) -> dict[int, str]:
        titles = {column.title: name for name, column in schema.columns.items()}
        mapping: dict[int, str] = {}
        seen: set[str] = set()
        issues: list[ExcelIssue] = []
        if len(header) > self.settings.max_columns:
            raise ExcelError(ExcelErrorCodes.LIMIT, "表头列数超过限制")
        for index, cell in enumerate(header, start=1):
            if cell.value is None:
                continue
            title = str(cell.value).strip()
            if not title:
                continue
            if cell.data_type in ("f", "e") or len(title) > self.settings.max_cell_text_length:
                issues.append(ExcelIssue(1, index, "", "表头类型或长度无效"))
            elif title in seen:
                issues.append(ExcelIssue(1, index, titles.get(title, ""), "表头重复"))
            elif title in titles:
                mapping[index] = titles[title]
            elif self.settings.unknown_columns == "reject":
                issues.append(ExcelIssue(1, index, "", "表头未在模型声明"))
            seen.add(title)
        for name, field in schema.model.model_fields.items():
            if field.is_required() and name not in mapping.values():
                issues.append(ExcelIssue(1, None, name, "缺少必需字段的表头"))
        if not mapping and not issues:
            issues.append(ExcelIssue(1, None, "", "没有匹配的表头"))
        if issues:
            self._raise_issues(issues[: self.settings.max_errors], None)
        return mapping

    @staticmethod
    def _raise_issues(issues: list[ExcelIssue], cause: Exception | None) -> None:
        raise ExcelError(
            ExcelErrorCodes.VALIDATION, "Excel 导入校验失败", issues=issues, cause=cause
        )
